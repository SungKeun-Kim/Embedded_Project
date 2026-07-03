"""
Generate full Korean/English LCD-panel user manuals.

This version is structured like the original generator manual:
summary, specification, cautions, part names, connection method,
operation preparation, detailed operating method, 25-pin D-SUB wiring,
error display, communication specification, initial settings, maintenance
and A/S.
"""
from __future__ import annotations

import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


DOCS = Path(__file__).resolve().parent
ASSET_PANEL = Path(
    r"C:\Users\USER\.cursor\projects\d-01-AllData-06-Embedded-Proj-STM32-Cop-ms-main-board-st"
    r"\assets\c__Users_USER_AppData_Roaming_Cursor_User_workspaceStorage_b6e626343dbe3874e127a8de46edff03_images_LCD_Display-238d1572-4b75-49c1-95e2-b09f6b7f8522.png"
)
BASE = "MANUAL_CSF_JET_LCD_Panel_v1"


def load_font(name: str, size: int):
    for suffix in ("ttf", "TTF"):
        path = f"C:/Windows/Fonts/{name}.{suffix}"
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


FONT_TITLE = load_font("malgunbd", 38)
FONT_HEAD = load_font("malgunbd", 28)
FONT = load_font("malgun", 23)
FONT_SMALL = load_font("malgun", 19)
FONT_TINY = load_font("malgun", 15)
FONT_LCD = load_font("consola", 42)
FONT_LCD_SMALL = load_font("consola", 34)


def md_table(headers, rows) -> str:
    out = "| " + " | ".join(headers) + " |\n"
    out += "| " + " | ".join(["---"] * len(headers)) + " |\n"
    for row in rows:
        out += "| " + " | ".join(str(v) for v in row) + " |\n"
    return out


def bullet(items) -> str:
    return "\n".join(f"- {item}" for item in items) + "\n"


def numbered(items) -> str:
    return "\n".join(f"{i + 1}. {item}" for i, item in enumerate(items)) + "\n"


def draw_center(draw: ImageDraw.ImageDraw, box, text: str, font, fill=(30, 30, 30), spacing=5):
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    x = box[0] + ((box[2] - box[0]) - (bbox[2] - bbox[0])) / 2
    y = box[1] + ((box[3] - box[1]) - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=font, fill=fill, spacing=spacing, align="center")


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(35, 85, 135), width=5):
    draw.line((start, end), fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        s = 16 if x2 >= x1 else -16
        draw.polygon([(x2, y2), (x2 - s, y2 - 10), (x2 - s, y2 + 10)], fill=fill)
    else:
        s = 16 if y2 >= y1 else -16
        draw.polygon([(x2, y2), (x2 - 10, y2 - s), (x2 + 10, y2 - s)], fill=fill)


def canvas(path: Path, title: str, size=(1600, 1050)):
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((15, 15, size[0] - 15, size[1] - 15), radius=24, outline=(30, 75, 120), width=4)
    draw.text((50, 35), title, font=FONT_TITLE, fill=(25, 65, 110))
    return img, draw


def save_actual_panel_image() -> Path:
    out = DOCS / "lcd_panel_actual.png"
    if ASSET_PANEL.exists():
        img = Image.open(ASSET_PANEL).convert("RGB")
        # Keep the real PCB image large enough for DOCX/PDF readability.
        w, h = img.size
        scale = 1500 / max(w, 1)
        img = img.resize((1500, int(h * scale)), Image.Resampling.LANCZOS)
        img.save(out)
    elif not out.exists():
        Image.new("RGB", (1500, 720), (240, 240, 240)).save(out)
    return out


def make_screen_examples(lang: str) -> Path:
    kr = lang == "KR"
    path = DOCS / f"lcd_screen_examples_large_{lang}.png"
    title = "LCD 화면 예시 (확대)" if kr else "LCD Screen Examples (Large Text)"
    img, d = canvas(path, title, (1600, 1200))
    examples = [
        ("기본 정지" if kr else "Basic Stop", "0.00W     0000kHz", "NORMAL STOP PWR1"),
        ("FREQ 채널" if kr else "FREQ Channel", "01CH      0400kHz", "CH01 E0000 A0000"),
        ("FREQ 편집" if kr else "FREQ Edit", "START(3SEC):AUTO", "FREQ:0406kHz"),
        ("BIT 편집" if kr else "BIT Edit", "START(3SEC):AUTO", "BIT:0010 03/16"),
        ("POWER 편집" if kr else "POWER Edit", "P-SET P1", "DEF:0.50W"),
        ("알람" if kr else "Alarm", "   ALARM STOP", "ERR1 LOW ALARM"),
    ]
    for i, (cap, line1, line2) in enumerate(examples):
        col = i % 2
        row = i // 2
        x = 85 + col * 745
        y = 145 + row * 320
        d.text((x, y), cap, font=FONT_HEAD, fill=(25, 65, 110))
        d.rounded_rectangle((x, y + 55, x + 650, y + 245), radius=16, fill=(26, 52, 36), outline=(60, 135, 85), width=5)
        d.text((x + 45, y + 90), line1, font=FONT_LCD, fill=(180, 255, 180))
        d.text((x + 45, y + 155), line2, font=FONT_LCD, fill=(180, 255, 180))
    img.save(path)
    return path


def make_menu_flow(lang: str) -> Path:
    kr = lang == "KR"
    path = DOCS / f"lcd_menu_flow_full_{lang}.png"
    img, d = canvas(path, "LCD 메뉴 설정 흐름도" if kr else "LCD Menu Setting Flow", (1600, 1150))
    texts = {
        "select": "SELECT 화면\nMODE / FREQ CH / 8 POWER\nSTART/STOP: RUN/STOP" if kr else "SELECT Screen\nMODE / FREQ CH / 8 POWER\nSTART/STOP: RUN/STOP",
        "setting": "SETTING MODE\n1.FREQ SET\n2.POWER SET\n3.EXT:RS485" if kr else "SETTING MODE\n1.FREQ SET\n2.POWER SET\n3.EXT:RS485",
        "freq": "FREQ SET\nCH -> FREQ -> BIT\nSET: 저장\nSTART 3초: Auto Tuning" if kr else "FREQ SET\nCH -> FREQ -> BIT\nSET: Save\nSTART 3 sec: Auto Tuning",
        "power": "POWER SET\nCH -> LOW -> HIGH -> DEF\nSET: 저장" if kr else "POWER SET\nCH -> LOW -> HIGH -> DEF\nSET: Save",
        "rs485": "EXT:RS485\nBAUD -> ADDR -> TERM -> PARITY\nSET: 저장" if kr else "EXT:RS485\nBAUD -> ADDR -> TERM -> PARITY\nSET: Save",
    }
    boxes = [
        (90, 160, 500, 330, texts["select"]),
        (595, 160, 1005, 330, texts["setting"]),
        (90, 620, 500, 840, texts["freq"]),
        (595, 620, 1005, 840, texts["power"]),
        (1090, 620, 1500, 840, texts["rs485"]),
    ]
    for x1, y1, x2, y2, text in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=(240, 248, 255), outline=(35, 85, 135), width=4)
        draw_center(d, (x1, y1, x2, y2), text, FONT)
    arrow(d, (500, 245), (595, 245))
    d.text((515, 200), "MODE LONG" if not kr else "MODE 길게", font=FONT_SMALL, fill=(35, 85, 135))
    arrow(d, (800, 330), (295, 620))
    arrow(d, (800, 330), (800, 620))
    arrow(d, (800, 330), (1295, 620))
    d.text((520, 455), "UP/DOWN SELECT -> SET" if not kr else "UP/DOWN 선택 -> SET", font=FONT, fill=(35, 85, 135))
    d.text((110, 950), "공통: MODE 짧게=필드 이동 / UP-DOWN=값 변경 / SET=저장 / SET 길게=선택 화면 복귀" if kr else "Common: MODE short=next field / UP-DOWN=change value / SET=save / SET long=return to select screen", font=FONT, fill=(45, 45, 45))
    img.save(path)
    return path


def make_dsub_pinout(lang: str) -> Path:
    kr = lang == "KR"
    path = DOCS / f"dsub25_pinout_{lang}.png"
    title = "25pin D-SUB 커넥션" if kr else "25-pin D-SUB Connection"
    img, d = canvas(path, title, (1600, 1120))
    d.rounded_rectangle((145, 190, 1455, 590), radius=70, fill=(238, 241, 244), outline=(80, 80, 80), width=5)
    d.rounded_rectangle((215, 260, 1385, 520), radius=45, fill=(220, 225, 230), outline=(80, 80, 80), width=4)

    top = list(range(1, 14))
    bottom = list(range(14, 26))
    pin_info = {
        1: ("LOW", (240, 210, 210)), 14: ("LOW", (240, 210, 210)),
        2: ("GO", (210, 240, 210)), 15: ("GO", (210, 240, 210)),
        3: ("HIGH", (255, 230, 190)), 16: ("HIGH", (255, 230, 190)),
        6: ("TRANS", (225, 220, 255)), 19: ("TRANS", (225, 220, 255)),
        7: ("OPER", (220, 235, 255)), 20: ("OPER", (220, 235, 255)),
        9: ("COM", (235, 235, 235)), 22: ("REMOTE", (255, 245, 190)),
        10: ("BCD3", (235, 255, 235)), 23: ("BCD2", (235, 255, 235)), 11: ("BCD1", (235, 255, 235)),
        13: ("SENSOR", (255, 220, 220)), 25: ("SENSOR", (255, 220, 220)),
    }
    for i, pin in enumerate(top):
        x = 275 + i * 86
        y = 330
        label, color = pin_info.get(pin, ("N.C.", (250, 250, 250)))
        d.ellipse((x, y, x + 45, y + 45), fill=color, outline=(70, 70, 70), width=2)
        draw_center(d, (x, y, x + 45, y + 45), str(pin), FONT_TINY)
        d.text((x - 18, y - 36), label, font=FONT_TINY, fill=(40, 40, 40))
    for i, pin in enumerate(bottom):
        x = 315 + i * 86
        y = 430
        label, color = pin_info.get(pin, ("N.C.", (250, 250, 250)))
        d.ellipse((x, y, x + 45, y + 45), fill=color, outline=(70, 70, 70), width=2)
        draw_center(d, (x, y, x + 45, y + 45), str(pin), FONT_TINY)
        d.text((x - 18, y + 55), label, font=FONT_TINY, fill=(40, 40, 40))

    notes = [
        ("SENSOR: 13-25 SHORT = normal" if not kr else "SENSOR: 13-25 SHORT = 정상"),
        ("REMOTE: 9(COM)-22 SHORT = RUN" if not kr else "REMOTE: 9(COM)-22 SHORT = RUN"),
        ("BCD1/2/3: 8 POWER external selection" if not kr else "BCD1/2/3: 8 POWER 외부 선택"),
        ("Relay outputs are dry contacts. Check external circuit ratings." if not kr else "알람 출력은 무전압 릴레이 접점이다. 외부 회로 정격을 확인한다."),
    ]
    for i, note in enumerate(notes):
        d.text((165, 675 + i * 55), note, font=FONT, fill=(50, 50, 50))
    img.save(path)
    return path


def make_connection_block(lang: str) -> Path:
    kr = lang == "KR"
    path = DOCS / f"connection_overview_{lang}.png"
    img, d = canvas(path, "후면 연결 구성" if kr else "Rear Connection Overview", (1600, 980))
    boxes = [
        (70, 170, 380, 330, "전원 입력\nAC / PE" if kr else "Power Input\nAC / PE"),
        (455, 170, 765, 330, "출력 커넥터\nNozzle / PZT" if kr else "Output Connector\nNozzle / PZT"),
        (840, 170, 1150, 330, "25pin D-SUB\nSENSOR/REMOTE/ALARM" if kr else "25-pin D-SUB\nSENSOR/REMOTE/ALARM"),
        (1225, 170, 1535, 330, "RS-485\nA/B + GND" if kr else "RS-485\nA/B + GND"),
        (455, 575, 765, 735, "보호접지\nPE 확인" if kr else "Protective Earth\nCheck PE"),
        (840, 575, 1150, 735, "PLC / PC\nModbus RTU" if kr else "PLC / PC\nModbus RTU"),
    ]
    for x1, y1, x2, y2, text in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=(244, 248, 252), outline=(35, 85, 135), width=4)
        draw_center(d, (x1, y1, x2, y2), text, FONT)
    arrow(d, (380, 250), (455, 250))
    arrow(d, (765, 250), (840, 250))
    arrow(d, (1150, 250), (1225, 250))
    arrow(d, (1000, 575), (1000, 330))
    d.text((105, 845), "전원, 출력, SENSOR, REMOTE, RS-485를 모두 확인한 후 운전한다." if kr else "Operate only after checking power, output, SENSOR, REMOTE and RS-485 wiring.", font=FONT, fill=(180, 45, 45))
    img.save(path)
    return path


def generate_figures(lang: str) -> dict[str, Path]:
    return {
        "actual_panel": save_actual_panel_image(),
        "screens": make_screen_examples(lang),
        "menu": make_menu_flow(lang),
        "dsub": make_dsub_pinout(lang),
        "connection": make_connection_block(lang),
    }


PINOUT_KR = [
    [1, "LOW ALARM A", "릴레이 접점", "Err1 발생 시 14번과 SHORT"],
    [14, "LOW ALARM B", "릴레이 접점", "Err1 LOW ALARM"],
    [2, "GO A", "릴레이 접점", "정상 출력 범위에서 15번과 SHORT"],
    [15, "GO B", "릴레이 접점", "정상 출력 범위"],
    [3, "HIGH ALARM A", "릴레이 접점", "Err2 발생 시 16번과 SHORT"],
    [16, "HIGH ALARM B", "릴레이 접점", "Err2 HIGH ALARM"],
    [6, "TRANSDUCER ALARM A", "릴레이 접점", "Err5 발생 시 19번과 SHORT"],
    [19, "TRANSDUCER ALARM B", "릴레이 접점", "진동자/부하 이상"],
    [7, "OPERATE A", "Fail-safe 릴레이", "정상 시 20번과 SHORT, 알람/전원 OFF 시 OPEN"],
    [20, "OPERATE B", "Fail-safe 릴레이", "통합 운전/알람 접점"],
    [9, "COM", "입력 공통", "REMOTE/BCD 입력 공통"],
    [22, "REMOTE", "접점 입력", "9번과 SHORT 시 RUN, OPEN 시 STOP"],
    [11, "BCD1", "접점 입력", "8 POWER 외부 선택 bit1"],
    [23, "BCD2", "접점 입력", "8 POWER 외부 선택 bit2"],
    [10, "BCD3", "접점 입력", "8 POWER 외부 선택 bit3"],
    [13, "SENSOR A", "안전 입력", "25번과 SHORT 정상"],
    [25, "SENSOR B", "안전 입력", "OPEN 시 Err7/Err8"],
    ["4,5,8,12,17,18,21,24", "N.C.", "예약", "사용하지 않음"],
    ["SHIELD", "Connector shell", "쉴드", "커넥터 바디 접지/쉴드"],
]

PINOUT_EN = [
    [1, "LOW ALARM A", "Relay contact", "SHORT to pin 14 when Err1 occurs"],
    [14, "LOW ALARM B", "Relay contact", "Err1 LOW ALARM"],
    [2, "GO A", "Relay contact", "SHORT to pin 15 in normal output range"],
    [15, "GO B", "Relay contact", "Normal output range"],
    [3, "HIGH ALARM A", "Relay contact", "SHORT to pin 16 when Err2 occurs"],
    [16, "HIGH ALARM B", "Relay contact", "Err2 HIGH ALARM"],
    [6, "TRANSDUCER ALARM A", "Relay contact", "SHORT to pin 19 when Err5 occurs"],
    [19, "TRANSDUCER ALARM B", "Relay contact", "Transducer/load alarm"],
    [7, "OPERATE A", "Fail-safe relay", "SHORT to pin 20 in normal state; OPEN at alarm/power-off"],
    [20, "OPERATE B", "Fail-safe relay", "Integrated operation/alarm contact"],
    [9, "COM", "Input common", "Common for REMOTE/BCD inputs"],
    [22, "REMOTE", "Contact input", "SHORT to pin 9 = RUN, OPEN = STOP"],
    [11, "BCD1", "Contact input", "8 POWER external select bit1"],
    [23, "BCD2", "Contact input", "8 POWER external select bit2"],
    [10, "BCD3", "Contact input", "8 POWER external select bit3"],
    [13, "SENSOR A", "Safety input", "SHORT to pin 25 is normal"],
    [25, "SENSOR B", "Safety input", "OPEN causes Err7/Err8"],
    ["4,5,8,12,17,18,21,24", "N.C.", "Reserved", "Do not use"],
    ["SHIELD", "Connector shell", "Shield", "Connector body shield/ground"],
]


def data(lang: str) -> dict:
    kr = lang == "KR"
    if kr:
        return {
            "version": "v4.0-KR",
            "title": "LCD 패널 사용자 운전 매뉴얼",
            "purpose": "본 매뉴얼은 사용자가 장비 설치, 배선, LCD 메뉴 설정, NORMAL/REMOTE/EXT 운전, 알람 조치, 유지보수를 수행할 수 있도록 작성한 최종 사용자 문서이다.",
            "summary": [
                ["발진 시작/정지", "SELECT 화면에서 START/STOP"],
                ["항목 이동", "MODE 짧게: MODE -> FREQ CH -> 8 POWER"],
                ["설정 진입", "SELECT 화면에서 MODE 길게"],
                ["값 변경", "UP/DOWN, 길게 누르면 빠른 반복"],
                ["저장", "SET 짧게, 성공 시 비프 2회"],
                ["선택 화면 복귀", "설정 화면에서 SET 약 2초"],
                ["Auto Tuning", "FREQ SET에서 START/STOP 약 3초"],
                ["알람 해제", "원인 제거 후 알람 화면에서 SET"],
            ],
            "spec": [
                [1, "모델명", "CSF JET Multi Megasonic Generator"],
                [2, "표시 장치", "LCD1602, 2행 x 16문자"],
                [3, "버튼", "START/STOP, MODE, UP, DOWN, SET"],
                [4, "상태 표시", "NORMAL, H/L SET, 8 POWER, REMOTE, EXT, TX/RX LED"],
                [5, "동작 주파수", "10CH, 400kHz ~ 2200kHz"],
                [6, "주파수 설정", "수동 설정 및 Auto Tuning"],
                [7, "LC BIT", "16단계, BIT 0000 ~ 1111"],
                [8, "출력 설정", "8 POWER 프로파일, LOW/HIGH/DEF"],
                [9, "출력 범위", "0.20W ~ 5.00W"],
                [10, "제어 커넥터", "25pin D-SUB"],
                [11, "통신", "RS-485 Half-Duplex, Modbus RTU"],
                [12, "설정 저장", "내부 FLASH 저장, 전원 OFF 후 유지"],
            ],
            "install": [
                "진동 영향을 받지 않는 안정된 위치에 설치한다.",
                "산성/알칼리성 가스, 전도성 먼지, 유증기, 고온 환경을 피한다.",
                "방열을 위해 전후좌우 통풍 공간을 확보한다.",
                "보호접지(PE)를 먼저 연결한 후 전원을 인가한다.",
                "지정된 노즐/트랜스듀서와 케이블을 사용한다.",
                "출력 케이블 종류나 길이를 임의로 변경하지 않는다.",
                "배선 작업 전 반드시 전원을 끈다.",
            ],
            "use": [
                "출력 커넥터가 빠진 상태에서 운전하지 않는다.",
                "SENSOR가 OPEN이면 발진이 금지되거나 정지된다.",
                "알람 발생 후 원인 확인 없이 반복 재기동하지 않는다.",
                "설치 또는 노즐 교체 후에는 낮은 출력부터 단계적으로 운전한다.",
                "REMOTE ON/OFF는 충분한 간격을 두고 수행한다.",
                "통신 운전 시 주소, 속도, 패리티를 마스터와 일치시킨다.",
            ],
            "parts": [
                ["LCD", "출력, 주파수, 운전 모드, 설정 메뉴, 알람 표시"],
                ["NORMAL LED", "NORMAL 운전 모드 표시"],
                ["H/L SET LED", "알람 상/하한 설정 관련 표시"],
                ["8 POWER LED", "8 POWER 선택/설정 표시"],
                ["REMOTE LED", "REMOTE 모드 표시"],
                ["EXT LED", "EXT 통신 모드 표시"],
                ["TX/RX LED", "RS-485 송수신 상태 표시"],
                ["START/STOP", "발진 시작/정지, FREQ SET에서 길게 누르면 Auto Tuning"],
                ["MODE", "선택 항목/편집 필드 이동, 길게 누르면 설정 진입 또는 상위 복귀"],
                ["UP/DOWN", "값 증가/감소, 길게 누르면 빠른 반복"],
                ["SET", "저장/확정, 설정 화면에서 길게 누르면 SELECT 화면 복귀"],
            ],
            "pinout": PINOUT_KR,
            "bcd": [["BCD3", "BCD2", "BCD1", "선택 POWER"], ["0", "0", "0", "PWR1"], ["0", "0", "1", "PWR2"], ["0", "1", "0", "PWR3"], ["0", "1", "1", "PWR4"], ["1", "0", "0", "PWR5"], ["1", "0", "1", "PWR6"], ["1", "1", "0", "PWR7"], ["1", "1", "1", "PWR8"]],
            "errors": [
                ["Err1", "LOW ALARM", "출력이 LOW 설정값보다 낮은 상태 지속", "출력 정지", "부하/케이블/LOW 설정 확인 후 SET"],
                ["Err2", "HIGH ALARM", "출력이 HIGH 설정값보다 높은 상태 지속", "출력 정지", "출력 설정/부하/HIGH 설정 확인 후 SET"],
                ["Err5", "TRANSDUCER ALARM", "진동자 또는 부하 이상", "출력 정지", "노즐/트랜스듀서/케이블 확인"],
                ["Err6", "SETTING ERROR", "허용 범위 밖 설정", "조건부", "허용 범위로 재입력"],
                ["Err7", "SENSOR OFF", "정지 중 SENSOR OPEN", "출력 금지", "전원 OFF -> SENSOR SHORT 확인 -> 전원 ON"],
                ["Err8", "SENSOR RUN", "운전 중 SENSOR OPEN", "출력 정지", "전원 OFF -> SENSOR SHORT 확인 -> 전원 ON"],
            ],
        }
    return {
        "version": "v4.0-EN",
        "title": "LCD Panel User Operation Manual",
        "purpose": "This final user manual explains installation, wiring, LCD menu setup, NORMAL/REMOTE/EXT operation, alarm recovery, maintenance and service information.",
        "summary": [
            ["Run/Stop", "Press START/STOP on the SELECT screen"],
            ["Move item", "Short press MODE: MODE -> FREQ CH -> 8 POWER"],
            ["Enter setting", "Long press MODE on the SELECT screen"],
            ["Change value", "UP/DOWN; long press repeats quickly"],
            ["Save", "Short press SET; two beeps indicate success"],
            ["Return to select screen", "Long press SET for about 2 seconds in setting screens"],
            ["Auto Tuning", "Long press START/STOP for about 3 seconds in FREQ SET"],
            ["Clear alarm", "Remove cause, then press SET on the alarm screen"],
        ],
        "spec": [
            [1, "Model name", "CSF JET Multi Megasonic Generator"],
            [2, "Display", "LCD1602, 2 lines x 16 characters"],
            [3, "Buttons", "START/STOP, MODE, UP, DOWN, SET"],
            [4, "Indicators", "NORMAL, H/L SET, 8 POWER, REMOTE, EXT, TX/RX LEDs"],
            [5, "Operating frequency", "10CH, 400kHz to 2200kHz"],
            [6, "Frequency setting", "Manual setting and Auto Tuning"],
            [7, "LC BIT", "16 steps, BIT 0000 to 1111"],
            [8, "Output setting", "8 POWER profiles, LOW/HIGH/DEF"],
            [9, "Output range", "0.20W to 5.00W"],
            [10, "Control connector", "25-pin D-SUB"],
            [11, "Communication", "RS-485 Half-Duplex, Modbus RTU"],
            [12, "Setting storage", "Internal FLASH, retained after power-off"],
        ],
        "install": [
            "Install the generator on a stable support where vibration does not affect the unit.",
            "Avoid acidic/alkaline gas, conductive dust, oil mist and high-temperature environments.",
            "Secure ventilation space around the unit.",
            "Connect protective earth before applying power.",
            "Use the assigned nozzle/transducer and cable.",
            "Do not change output cable type or length without confirmation.",
            "Turn off power before wiring work.",
        ],
        "use": [
            "Do not operate while the output connector is disconnected.",
            "If SENSOR is open, oscillation is inhibited or stopped.",
            "Do not repeatedly restart after an alarm without checking the cause.",
            "After installation or nozzle replacement, start at low output and increase gradually.",
            "Keep sufficient interval between REMOTE ON and OFF.",
            "For communication operation, match address, baudrate and parity with the master.",
        ],
        "parts": [
            ["LCD", "Displays output, frequency, operation mode, setting menu and alarms"],
            ["NORMAL LED", "Indicates NORMAL operation mode"],
            ["H/L SET LED", "Indicates high/low alarm setting related status"],
            ["8 POWER LED", "Indicates 8 POWER selection/setting"],
            ["REMOTE LED", "Indicates REMOTE mode"],
            ["EXT LED", "Indicates EXT communication mode"],
            ["TX/RX LED", "Indicates RS-485 transmit/receive activity"],
            ["START/STOP", "Starts/stops oscillation; long press in FREQ SET starts Auto Tuning"],
            ["MODE", "Moves selected item/edit field; long press enters setting or returns to upper menu"],
            ["UP/DOWN", "Increases/decreases values; long press repeats quickly"],
            ["SET", "Saves/confirms values; long press returns to SELECT screen in setting screens"],
        ],
        "pinout": PINOUT_EN,
        "bcd": [["BCD3", "BCD2", "BCD1", "Selected POWER"], ["0", "0", "0", "PWR1"], ["0", "0", "1", "PWR2"], ["0", "1", "0", "PWR3"], ["0", "1", "1", "PWR4"], ["1", "0", "0", "PWR5"], ["1", "0", "1", "PWR6"], ["1", "1", "0", "PWR7"], ["1", "1", "1", "PWR8"]],
        "errors": [
            ["Err1", "LOW ALARM", "Output remains below LOW setting", "Stop", "Check load/cable/LOW setting, then press SET"],
            ["Err2", "HIGH ALARM", "Output remains above HIGH setting", "Stop", "Check output setting/load/HIGH setting, then press SET"],
            ["Err5", "TRANSDUCER ALARM", "Transducer or load abnormal", "Stop", "Check nozzle/transducer/cable"],
            ["Err6", "SETTING ERROR", "Setting outside allowed range", "Conditional", "Enter a value within the allowed range"],
            ["Err7", "SENSOR OFF", "SENSOR open while stopped", "Output inhibited", "Power OFF -> check SENSOR short -> power ON"],
            ["Err8", "SENSOR RUN", "SENSOR open while running", "Stop", "Power OFF -> check SENSOR short -> power ON"],
        ],
    }


def common_tables(lang: str):
    kr = lang == "KR"
    freq = [[f"CH{i:02d}", f"{400 + (i - 1) * 200}kHz", f"{'주파수 프로파일' if kr else 'Frequency profile'} {i}"] for i in range(1, 11)]
    power = [["CH", "POWER 프로파일 번호" if kr else "POWER profile number", "PWR1~PWR8"],
             ["LOW", "저출력 알람 기준" if kr else "Low alarm threshold", "0.10W~0.95W"],
             ["HIGH", "고출력 알람 기준" if kr else "High alarm threshold", "0.20W~5.00W"],
             ["DEF", "기본 출력" if kr else "Default output", "0.20W~5.00W"]]
    rs485 = [["BAUD", "9600/19200/38400/115200", "통신 속도" if kr else "Communication speed"],
             ["ADDR", "1~247", "슬레이브 주소" if kr else "Slave address"],
             ["TERM", "OFF/ON", "종단 설정 기록, 하드웨어 제어 예약" if kr else "Termination setting record; hardware control reserved"],
             ["PARITY", "NONE-8-1/EVEN-8-1", "패널 설정 패리티" if kr else "Panel parity setting"]]
    regs = [["0x0000", "R/W", "RUN 제어" if kr else "RUN control", "0=STOP, 1=RUN"],
            ["0x0001", "R/W", "주파수" if kr else "Frequency", "x0.1kHz"],
            ["0x0002", "R/W", "Duty", "x0.1%"],
            ["0x0006", "R/W", "운전 모드" if kr else "Operation mode", "0=NORMAL,1=REMOTE,2=EXT"],
            ["0x0010", "R/W", "슬레이브 주소" if kr else "Slave address", "1~247"],
            ["0x0011", "R/W", "Baud 인덱스" if kr else "Baud index", "0~3"],
            ["0x0012", "R/W", "Parity", "0=None,1=Even,2=Odd"]]
    return freq, power, rs485, regs


def build_markdown(lang: str, figures: dict[str, Path]) -> str:
    kr = lang == "KR"
    d = data(lang)
    freq, power, rs485, regs = common_tables(lang)
    t = {
        "item": "항목" if kr else "Item",
        "operation": "조작" if kr else "Operation",
        "spec": "사양" if kr else "Specification",
        "part": "부품" if kr else "Part",
        "function": "기능" if kr else "Function",
        "pin": "핀" if kr else "Pin",
        "signal": "신호" if kr else "Signal",
        "type": "구분" if kr else "Type",
        "desc": "설명" if kr else "Description",
    }
    out = f"# CSF JET Multi Megasonic Generator\n\n## {d['title']}\n\n"
    out += f"- {'문서 버전' if kr else 'Document Version'}: {d['version']}\n- {'개정일' if kr else 'Revision Date'}: 2026-05-19\n- {'적용 장비' if kr else 'Applicable Product'}: CSF JET Multi Megasonic Generator (LCD Panel Type)\n- {'목적' if kr else 'Purpose'}: {d['purpose']}\n\n---\n\n"
    out += f"## < OPERATING SUMMARY >\n\n{md_table([t['item'], t['operation']], d['summary'])}\n"
    out += ("비프음: 일반 버튼 1회, 길게 누름 인식 긴 1회, 저장 성공 2회, 저장 실패 3회, 알람 긴 1회.\n\n" if kr else "Buzzer guide: normal key 1 beep, long-press recognition 1 longer beep, save success 2 beeps, save failure 3 beeps, alarm 1 long beep.\n\n")
    out += f"## 1. Specification\n\n{md_table(['No.', t['item'], t['spec']], d['spec'])}\n"
    out += "## 2. Caution on Product Use\n\n"
    out += f"### {'2.1 설치 주의사항' if kr else '2.1 Caution on Installation'}\n\n{numbered(d['install'])}\n"
    out += f"### {'2.2 사용 주의사항' if kr else '2.2 Caution on Use'}\n\n{numbered(d['use'])}\n"
    out += f"## 3. Name and Function of Each Part\n\n### {'3.1 실제 LCD 조작 패널' if kr else '3.1 Actual LCD Operation Panel'}\n\n![panel](./{figures['actual_panel'].name})\n\n"
    out += f"### {'3.2 LCD 화면 예시' if kr else '3.2 LCD Screen Examples'}\n\n![screens](./{figures['screens'].name})\n\n"
    out += md_table([t["part"], t["function"]], d["parts"]) + "\n"
    out += f"## 4. Connecting Method\n\n![connection](./{figures['connection'].name})\n\n"
    out += ("전원, 출력, SENSOR, REMOTE, ALARM, RS-485를 모두 확인한 후 운전한다.\n\n" if kr else "Operate only after checking power, output, SENSOR, REMOTE, ALARM and RS-485 wiring.\n\n")
    out += f"### {'4.1 25pin D-SUB Control Connector' if kr else '4.1 25-pin D-SUB Control Connector'}\n\n![dsub](./{figures['dsub'].name})\n\n"
    out += md_table([t["pin"], t["signal"], t["type"], t["desc"]], d["pinout"]) + "\n"
    out += f"### {'4.2 REMOTE / SENSOR / ALARM 주의사항' if kr else '4.2 REMOTE / SENSOR / ALARM Notes'}\n\n"
    out += bullet([
        "SENSOR는 정상 상태에서 SHORT이며 OPEN 시 Err7/Err8이 발생한다." if kr else "SENSOR is normally SHORT. OPEN causes Err7/Err8.",
        "REMOTE는 COM과 REMOTE 입력이 SHORT이면 RUN, OPEN이면 STOP이다." if kr else "REMOTE runs when COM and REMOTE input are SHORT, and stops when OPEN.",
        "알람 출력은 무전압 릴레이 접점이므로 외부 전원/부하 정격을 확인한다." if kr else "Alarm outputs are dry relay contacts. Check external power/load rating.",
        "OPERATE는 Fail-safe 접점으로 정상 시 SHORT, 알람 또는 전원 OFF 시 OPEN이다." if kr else "OPERATE is fail-safe: SHORT in normal state, OPEN at alarm or power-off.",
    ])
    out += f"### {'4.3 8 POWER 외부 선택' if kr else '4.3 8 POWER External Selection'}\n\n{md_table(d['bcd'][0], d['bcd'][1:])}\n"
    out += f"## 5. Operation Ready\n\n{numbered(['공정 조건에 따라 노즐에 DI water 또는 액체가 공급되는지 확인한다.' if kr else 'Confirm DI water or process liquid is supplied to the nozzle when required.', '출력 케이블과 25pin D-SUB 배선을 확인한다.' if kr else 'Check output cable and 25-pin D-SUB wiring.', 'SENSOR가 정상 SHORT 상태인지 확인한다.' if kr else 'Confirm SENSOR is in normal SHORT state.', '전원 ON 후 LCD와 LED 표시를 확인한다.' if kr else 'Turn on power and check LCD/LED indication.', '낮은 출력으로 최초 운전을 시작한다.' if kr else 'Start the first operation at low output.'])}\n"
    out += f"## 6. Operating Method\n\n### {'6.1 메뉴 구조' if kr else '6.1 Menu Structure'}\n\n![menu](./{figures['menu'].name})\n\n"
    out += f"### {'6.2 NORMAL 운전' if kr else '6.2 NORMAL Operation'}\n\n{numbered(['전원 ON 후 NORMAL LED를 확인한다.' if kr else 'Turn on power and check the NORMAL LED.', 'MODE 항목에서 NORMAL을 선택하고 SET 저장한다.' if kr else 'Select NORMAL in the MODE item and save with SET.', 'FREQ CH와 8 POWER를 선택한다.' if kr else 'Select FREQ CH and 8 POWER.', 'START/STOP을 눌러 발진을 시작한다.' if kr else 'Press START/STOP to start oscillation.', 'LCD의 RUN 표시와 출력값을 확인한다.' if kr else 'Check RUN indication and output value on the LCD.', '정지하려면 START/STOP을 다시 누른다.' if kr else 'Press START/STOP again to stop.'])}\n"
    out += f"### {'6.3 FREQ 채널 선택' if kr else '6.3 FREQ Channel Selection'}\n\n{md_table(['채널' if kr else 'Channel', '기본 주파수' if kr else 'Default Frequency', '설명' if kr else 'Description'], freq)}\n"
    out += ("SELECT 화면에서 MODE로 FREQ CH 항목을 선택하고 UP/DOWN으로 채널을 고른 뒤 SET을 누른다.\n\n" if kr else "On the SELECT screen, use MODE to select FREQ CH, choose a channel with UP/DOWN, then press SET.\n\n")
    out += f"### {'6.4 FREQ SET: 수동 주파수/BIT 설정' if kr else '6.4 FREQ SET: Manual Frequency/BIT Setting'}\n\n"
    out += numbered([
        "MODE 길게 -> SETTING MODE 진입." if kr else "Long press MODE to enter SETTING MODE.",
        "1.FREQ SET 선택 후 SET." if kr else "Select 1.FREQ SET and press SET.",
        "MODE 짧게로 CH -> FREQ -> BIT 필드를 이동한다." if kr else "Short press MODE to move CH -> FREQ -> BIT fields.",
        "깜빡이는 값은 UP/DOWN으로 변경한다." if kr else "Change the blinking value with UP/DOWN.",
        "SET을 눌러 현재 CH의 주파수와 BIT를 저장한다." if kr else "Press SET to save the frequency and BIT of the current channel.",
        "저장 성공 시 비프음이 2회 울린다." if kr else "Two beeps indicate successful save.",
    ]) + "\n"
    out += ("BIT는 LC 릴레이 조합이며 0000~1111의 16단계이다. E는 저장된 기대 BIT, A는 현재 실제 적용 BIT이다.\n\n" if kr else "BIT is the LC relay combination, 16 steps from 0000 to 1111. E is the saved expected BIT; A is the currently applied actual BIT.\n\n")
    out += f"### {'6.5 Auto Tuning' if kr else '6.5 Auto Tuning'}\n\n"
    out += numbered([
        "FREQ SET에서 대상 CH를 선택한다." if kr else "Select the target channel in FREQ SET.",
        "START/STOP을 약 3초 누른다." if kr else "Press START/STOP for about 3 seconds.",
        "LCD에 튜닝 진행 상태가 표시된다." if kr else "The LCD shows tuning progress.",
        "AUTO TUNE DONE 표시 후 필요한 경우 SET 저장한다." if kr else "When AUTO TUNE DONE appears, save with SET if needed.",
        "NO DETECTED 또는 TUNE TIMEOUT이면 노즐/케이블/부하를 점검한다." if kr else "If NO DETECTED or TUNE TIMEOUT appears, check nozzle/cable/load.",
    ]) + "\n"
    out += f"### {'6.6 POWER SET' if kr else '6.6 POWER SET'}\n\n{md_table(['필드' if kr else 'Field', '의미' if kr else 'Meaning', '범위' if kr else 'Range'], power)}\n"
    out += ("MODE 길게 -> 2.POWER SET -> SET -> MODE로 CH/LOW/HIGH/DEF 이동 -> UP/DOWN 변경 -> SET 저장.\n\n" if kr else "Long press MODE -> 2.POWER SET -> SET -> use MODE to move CH/LOW/HIGH/DEF -> change with UP/DOWN -> save with SET.\n\n")
    out += f"### {'6.7 REMOTE 운전' if kr else '6.7 REMOTE Operation'}\n\n"
    out += numbered([
        "발진을 정지하고 전원을 끈다." if kr else "Stop oscillation and turn off power.",
        "25pin D-SUB의 REMOTE 접점을 배선한다." if kr else "Wire the REMOTE contact on the 25-pin D-SUB.",
        "전원 ON 후 MODE 항목에서 REMOTE를 선택하고 SET 저장한다." if kr else "Turn on power, select REMOTE in MODE item and save with SET.",
        "REMOTE 접점 SHORT 시 RUN, OPEN 시 STOP으로 동작한다." if kr else "REMOTE SHORT means RUN; OPEN means STOP.",
    ]) + "\n"
    out += f"### {'6.8 EXT 운전' if kr else '6.8 EXT Operation'}\n\n"
    out += numbered([
        "MODE 항목에서 EXT를 선택하고 SET 저장한다." if kr else "Select EXT in MODE item and save with SET.",
        "SETTING MODE -> 3.EXT:RS485로 진입한다." if kr else "Enter SETTING MODE -> 3.EXT:RS485.",
        "BAUD, ADDR, TERM, PARITY를 설정한다." if kr else "Set BAUD, ADDR, TERM and PARITY.",
        "마스터 장비에서 동일한 통신 조건으로 접속한다." if kr else "Connect from the master using the same communication settings.",
    ]) + "\n" + md_table(["필드" if kr else "Field", "설정값" if kr else "Setting", "설명" if kr else "Description"], rs485) + "\n"
    out += f"## 7. ERROR Display\n\n```text\n   ALARM STOP\nERR1 LOW ALARM\n```\n\n{md_table(['코드' if kr else 'Code', '내용' if kr else 'Contents', '발생 조건' if kr else 'Condition', '출력' if kr else 'Output', '조치' if kr else 'Action'], d['errors'])}\n"
    out += f"## 8. Communication Specification\n\n{md_table(['항목' if kr else 'Item', '사양' if kr else 'Specification'], [['인터페이스' if kr else 'Interface', 'RS-485 Half-Duplex'], ['프로토콜' if kr else 'Protocol', 'Modbus RTU'], ['기본 설정' if kr else 'Default setting', 'ADDR 1, 9600bps, NONE-8-1'], ['지원 속도' if kr else 'Supported baudrates', '9600/19200/38400/115200'], ['설정 메뉴' if kr else 'Setting menu', 'SETTING MODE -> 3.EXT:RS485'], ['기능 코드' if kr else 'Function codes', '0x03 / 0x06 / 0x10']])}\n"
    out += md_table(["주소" if kr else "Address", "R/W", "명칭" if kr else "Name", "값" if kr else "Value"], regs) + "\n"
    out += f"## 9. Initial Setting Value\n\n{md_table(['항목' if kr else 'Item', '초기값' if kr else 'Initial Value'], [['MODE', 'NORMAL'], ['FREQ CH', 'CH01 / 400kHz'], ['POWER', 'PWR1 / DEF 0.50W'], ['LOW ALARM', '0.10W'], ['HIGH ALARM', '5.00W'], ['RS-485', 'ADDR1 / 9600bps / NONE-8-1']])}\n"
    out += f"## 10. Setting Storage\n\n"
    out += ("SET으로 저장한 운전 모드, FREQ 채널/주파수/BIT, POWER 채널/LOW/HIGH/DEF, RS-485 설정은 내부 FLASH에 저장되어 전원 OFF 후에도 유지된다.\n\n" if kr else "Operating mode, FREQ channel/frequency/BIT, POWER channel/LOW/HIGH/DEF, and RS-485 settings saved by SET are stored in internal FLASH and retained after power-off.\n\n")
    out += f"## 11. Troubleshooting\n\n{md_table(['증상' if kr else 'Symptom', '원인 후보' if kr else 'Possible Cause', '조치' if kr else 'Action'], [['출력 없음' if kr else 'No output', 'STOP/알람/SENSOR OPEN' if kr else 'STOP/alarm/SENSOR open', 'LCD와 알람, SENSOR 확인' if kr else 'Check LCD, alarm and SENSOR'], ['전원 ON 시 Err7' if kr else 'Err7 at power ON', 'SENSOR OPEN', '전원 OFF 후 13-25 SHORT 확인' if kr else 'Power OFF and check pin 13-25 SHORT'], ['REMOTE 동작 안 됨' if kr else 'REMOTE does not work', '9-22 접점/모드 설정 오류' if kr else 'Pin 9-22 contact or mode setting error', 'REMOTE 배선과 모드 확인' if kr else 'Check REMOTE wiring and mode'], ['RS-485 응답 없음' if kr else 'No RS-485 response', 'A/B, 주소, 속도, 패리티 불일치' if kr else 'A/B, address, baudrate or parity mismatch', 'EXT:RS485 설정 확인' if kr else 'Check EXT:RS485 setting'], ['저장값 미복원' if kr else 'Saved value not restored', 'SET 미저장 또는 저장 실패' if kr else 'SET not pressed or save failed', 'SET 후 비프 2회 확인' if kr else 'Save again and confirm two beeps']])}\n"
    out += f"## 12. Maintenance\n\n" + bullet([
        "일일: 케이블/커넥터, LCD/LED, 알람 이력 확인" if kr else "Daily: check cables/connectors, LCD/LEDs and alarm history",
        "주간: 출력 재현성, REMOTE 접점, RS-485 응답 확인" if kr else "Weekly: check output repeatability, REMOTE contact and RS-485 response",
        "월간: 트랜스듀서/케이블 절연, 알람 기준, 중요 설정 백업" if kr else "Monthly: check transducer/cable insulation, alarm thresholds and setting backup",
    ]) + "\n"
    out += f"## 13. A/S\n\n"
    out += ("서비스 요청 시 모델명, 시리얼 번호, 알람 코드, 발생 조건, 설치 사진, 재현 절차를 함께 제공한다. 임의 개조, 오배선, 자연재해, 소모품은 보증 범위에서 제외될 수 있다.\n\n" if kr else "When requesting service, provide model name, serial number, alarm code, operating condition, installation photos and reproduction steps. Unauthorized modification, incorrect wiring, natural disasters and consumables may be excluded from warranty.\n\n")
    out += f"## 14. Revision History\n\n{md_table(['Version', 'Date', 'Description'], [[d['version'], '2026-05-19', 'LCD panel full user manual with 25-pin D-SUB, actual panel image, detailed menu settings, REMOTE/EXT operation and bilingual output']])}\n"
    return out


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(head)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = "맑은 고딕"
                r.font.size = Pt(8.5)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "맑은 고딕"
                    r.font.size = Pt(8.5)
    doc.add_paragraph()


def add_image(doc: Document, path: Path, caption: str, width=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9)


def add_numbered(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_bullets(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_docx(lang: str, figures: dict[str, Path], out_path: Path):
    kr = lang == "KR"
    d = data(lang)
    freq, power, rs485, regs = common_tables(lang)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "맑은 고딕"
    doc.styles["Normal"].font.size = Pt(9.5)
    for level in range(1, 4):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "맑은 고딕"
        st.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Instruction Manual")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("----- CSF JET Multi Megasonic Generator -----")
    run.bold = True
    run.font.size = Pt(16)
    info = doc.add_paragraph(f"{d['title']} / {d['version']} / 2026-05-19")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_image(doc, figures["actual_panel"], "Actual LCD operation panel", 6.7)
    doc.add_page_break()

    doc.add_heading("< OPERATING SUMMARY >", level=1)
    add_table(doc, ["항목" if kr else "Item", "조작" if kr else "Operation"], d["summary"])
    doc.add_heading("1. Specification", level=1)
    add_table(doc, ["No.", "항목" if kr else "Item", "사양" if kr else "Specification"], d["spec"])

    doc.add_heading("2. Caution on Product Use", level=1)
    doc.add_heading("2.1 설치 주의사항" if kr else "2.1 Caution on Installation", level=2)
    add_numbered(doc, d["install"])
    doc.add_heading("2.2 사용 주의사항" if kr else "2.2 Caution on Use", level=2)
    add_numbered(doc, d["use"])

    doc.add_heading("3. Name and Function of Each Part", level=1)
    add_image(doc, figures["screens"], "LCD screen examples", 6.7)
    add_table(doc, ["부품" if kr else "Part", "기능" if kr else "Function"], d["parts"])

    doc.add_heading("4. Connecting Method", level=1)
    add_image(doc, figures["connection"], "Rear connection overview", 6.7)
    doc.add_heading("4.1 25pin D-SUB Control Connector" if kr else "4.1 25-pin D-SUB Control Connector", level=2)
    add_image(doc, figures["dsub"], "25-pin D-SUB connection", 6.7)
    add_table(doc, ["핀" if kr else "Pin", "신호" if kr else "Signal", "구분" if kr else "Type", "설명" if kr else "Description"], d["pinout"])
    doc.add_heading("4.2 8 POWER External Selection", level=2)
    add_table(doc, d["bcd"][0], d["bcd"][1:])

    doc.add_heading("5. Operation Ready", level=1)
    add_numbered(doc, [
        "공정 조건에 따라 노즐에 DI water 또는 액체가 공급되는지 확인한다." if kr else "Confirm DI water or process liquid is supplied to the nozzle when required.",
        "출력 케이블과 25pin D-SUB 배선을 확인한다." if kr else "Check output cable and 25-pin D-SUB wiring.",
        "SENSOR가 정상 SHORT 상태인지 확인한다." if kr else "Confirm SENSOR is in normal SHORT state.",
        "전원 ON 후 LCD와 LED 표시를 확인한다." if kr else "Turn on power and check LCD/LED indication.",
        "낮은 출력으로 최초 운전을 시작한다." if kr else "Start the first operation at low output.",
    ])

    doc.add_heading("6. Operating Method", level=1)
    add_image(doc, figures["menu"], "LCD menu setting flow", 6.7)
    doc.add_heading("6.1 FREQ Channel", level=2)
    add_table(doc, ["채널" if kr else "Channel", "기본 주파수" if kr else "Default Frequency", "설명" if kr else "Description"], freq)
    doc.add_heading("6.2 POWER SET", level=2)
    add_table(doc, ["필드" if kr else "Field", "의미" if kr else "Meaning", "범위" if kr else "Range"], power)
    doc.add_heading("6.3 EXT:RS485", level=2)
    add_table(doc, ["필드" if kr else "Field", "설정값" if kr else "Setting", "설명" if kr else "Description"], rs485)
    add_bullets(doc, [
        "MODE 길게 -> SETTING MODE -> 항목 선택 -> SET 진입 -> MODE 필드 이동 -> UP/DOWN 변경 -> SET 저장" if kr else "MODE long -> SETTING MODE -> select item -> SET enter -> MODE next field -> UP/DOWN change -> SET save",
        "FREQ SET에서 START/STOP 약 3초: Auto Tuning" if kr else "START/STOP for about 3 seconds in FREQ SET: Auto Tuning",
        "저장 성공: 비프 2회, 저장 실패: 비프 3회" if kr else "Save success: 2 beeps, save failure: 3 beeps",
    ])

    doc.add_heading("7. ERROR Display", level=1)
    add_table(doc, ["코드" if kr else "Code", "내용" if kr else "Contents", "발생 조건" if kr else "Condition", "출력" if kr else "Output", "조치" if kr else "Action"], d["errors"])

    doc.add_heading("8. Communication Specification", level=1)
    add_table(doc, ["항목" if kr else "Item", "사양" if kr else "Specification"], [
        ["인터페이스" if kr else "Interface", "RS-485 Half-Duplex"],
        ["프로토콜" if kr else "Protocol", "Modbus RTU"],
        ["기본 설정" if kr else "Default setting", "ADDR 1, 9600bps, NONE-8-1"],
        ["지원 속도" if kr else "Supported baudrates", "9600/19200/38400/115200"],
        ["설정 메뉴" if kr else "Setting menu", "SETTING MODE -> 3.EXT:RS485"],
        ["기능 코드" if kr else "Function codes", "0x03 / 0x06 / 0x10"],
    ])
    add_table(doc, ["주소" if kr else "Address", "R/W", "명칭" if kr else "Name", "값" if kr else "Value"], regs)

    doc.add_heading("9. Initial Setting Value", level=1)
    add_table(doc, ["항목" if kr else "Item", "초기값" if kr else "Initial Value"], [["MODE", "NORMAL"], ["FREQ CH", "CH01 / 400kHz"], ["POWER", "PWR1 / DEF 0.50W"], ["LOW ALARM", "0.10W"], ["HIGH ALARM", "5.00W"], ["RS-485", "ADDR1 / 9600bps / NONE-8-1"]])

    doc.add_heading("10. Setting Storage", level=1)
    doc.add_paragraph("SET으로 저장한 운전 모드, FREQ 채널/주파수/BIT, POWER 채널/LOW/HIGH/DEF, RS-485 설정은 내부 FLASH에 저장되어 전원 OFF 후에도 유지된다." if kr else "Operating mode, FREQ channel/frequency/BIT, POWER channel/LOW/HIGH/DEF and RS-485 settings saved by SET are stored in internal FLASH and retained after power-off.")

    doc.add_heading("11. Troubleshooting", level=1)
    add_table(doc, ["증상" if kr else "Symptom", "원인 후보" if kr else "Possible Cause", "조치" if kr else "Action"], [
        ["출력 없음" if kr else "No output", "STOP/알람/SENSOR OPEN" if kr else "STOP/alarm/SENSOR open", "LCD와 알람, SENSOR 확인" if kr else "Check LCD, alarm and SENSOR"],
        ["전원 ON 시 Err7" if kr else "Err7 at power ON", "SENSOR OPEN", "전원 OFF 후 13-25 SHORT 확인" if kr else "Power OFF and check pin 13-25 SHORT"],
        ["REMOTE 동작 안 됨" if kr else "REMOTE does not work", "9-22 접점/모드 설정 오류" if kr else "Pin 9-22 contact or mode setting error", "REMOTE 배선과 모드 확인" if kr else "Check REMOTE wiring and mode"],
        ["RS-485 응답 없음" if kr else "No RS-485 response", "A/B, 주소, 속도, 패리티 불일치" if kr else "A/B, address, baudrate or parity mismatch", "EXT:RS485 설정 확인" if kr else "Check EXT:RS485 setting"],
    ])
    doc.add_heading("12. Maintenance", level=1)
    add_bullets(doc, [
        "일일: 케이블/커넥터, LCD/LED, 알람 이력 확인" if kr else "Daily: check cables/connectors, LCD/LEDs and alarm history",
        "주간: 출력 재현성, REMOTE 접점, RS-485 응답 확인" if kr else "Weekly: check output repeatability, REMOTE contact and RS-485 response",
        "월간: 트랜스듀서/케이블 절연, 알람 기준, 중요 설정 백업" if kr else "Monthly: check transducer/cable insulation, alarm thresholds and setting backup",
    ])
    doc.add_heading("13. A/S", level=1)
    doc.add_paragraph("서비스 요청 시 모델명, 시리얼 번호, 알람 코드, 발생 조건, 설치 사진, 재현 절차를 함께 제공한다." if kr else "When requesting service, provide model name, serial number, alarm code, operating condition, installation photos and reproduction steps.")
    doc.add_heading("14. Revision History", level=1)
    add_table(doc, ["Version", "Date", "Description"], [[d["version"], "2026-05-19", "Full LCD panel user manual with 25-pin D-SUB, actual panel image, detailed menu settings, REMOTE/EXT operation and bilingual output"]])
    doc.save(out_path)


def write_manual(lang: str, md_path: Path, docx_path: Path):
    figures = generate_figures(lang)
    md_path.write_text(build_markdown(lang, figures), encoding="utf-8")
    build_docx(lang, figures, docx_path)


def main():
    outputs = [
        ("KR", DOCS / f"{BASE}_KR.md", DOCS / f"{BASE}_KR.docx"),
        ("EN", DOCS / f"{BASE}_EN.md", DOCS / f"{BASE}_EN.docx"),
        ("KR", DOCS / f"{BASE}.md", DOCS / f"{BASE}.docx"),
    ]
    for lang, md_path, docx_path in outputs:
        write_manual(lang, md_path, docx_path)
        print(md_path)
        print(docx_path)
    # Compatibility copy for users who still look for the previous actual-image name.
    panel = DOCS / "lcd_panel_actual.png"
    if panel.exists():
        shutil.copyfile(panel, DOCS / "LCD_Display.png")


if __name__ == "__main__":
    main()
