"""
Generate Korean and English LCD-panel user manuals.

Outputs:
- MANUAL_CSF_JET_LCD_Panel_v1.md / .docx     (Korean default)
- MANUAL_CSF_JET_LCD_Panel_v1_KR.md / .docx  (Korean copy)
- MANUAL_CSF_JET_LCD_Panel_v1_EN.md / .docx  (English)
"""
from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


DOCS = Path(__file__).resolve().parent
BASE = "MANUAL_CSF_JET_LCD_Panel_v1"


def load_font(name: str, size: int):
    for path in (f"C:/Windows/Fonts/{name}.ttf", f"C:/Windows/Fonts/{name}.TTF"):
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


FONT_TITLE = load_font("malgunbd", 30)
FONT_HEAD = load_font("malgunbd", 22)
FONT = load_font("malgun", 18)
FONT_SMALL = load_font("malgun", 15)
FONT_LCD = load_font("consola", 22)


def draw_center(draw: ImageDraw.ImageDraw, box, text: str, font, fill=(0, 0, 0), spacing=4):
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    x = box[0] + ((box[2] - box[0]) - (bbox[2] - bbox[0])) / 2
    y = box[1] + ((box[3] - box[1]) - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=font, fill=fill, align="center", spacing=spacing)


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(35, 85, 135), width=4):
    draw.line((start, end), fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 13 if x2 >= x1 else -13
        draw.polygon([(x2, y2), (x2 - sign, y2 - 8), (x2 - sign, y2 + 8)], fill=fill)
    else:
        sign = 13 if y2 >= y1 else -13
        draw.polygon([(x2, y2), (x2 - 8, y2 - sign), (x2 + 8, y2 - sign)], fill=fill)


def canvas(filename: str, title: str, size=(1200, 760)):
    path = DOCS / filename
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((10, 10, size[0] - 10, size[1] - 10), radius=20, outline=(30, 75, 120), width=3)
    draw.text((35, 28), title, font=FONT_TITLE, fill=(25, 65, 110))
    return img, draw, path


def make_figures(lang: str):
    kr = lang == "KR"
    prefix = f"lcd_user_{lang.lower()}"
    text = {
        "panel_title": "LCD 조작 패널 구성" if kr else "LCD Operation Panel",
        "lcd": "LCD1602 표시창" if kr else "LCD1602 Display",
        "mode": "항목/필드" if kr else "Item/Field",
        "up": "증가" if kr else "Increase",
        "down": "감소" if kr else "Decrease",
        "set": "저장" if kr else "Save",
        "flow_title": "LCD 메뉴 흐름도" if kr else "LCD Menu Flow",
        "select": "선택 화면\nRUN/STOP, MODE,\nFREQ CH, PWR 선택" if kr else "SELECT Screen\nRUN/STOP, MODE,\nFREQ CH, PWR Select",
        "setting": "SETTING MODE\n1.FREQ SET\n2.POWER SET\n3.EXT:RS485",
        "freq": "FREQ SET\nCH → FREQ → BIT\nSET 저장\nSTART 3초: AUTO" if kr else "FREQ SET\nCH → FREQ → BIT\nSET Save\nSTART 3 sec: AUTO",
        "power": "POWER SET\nCH → LOW → HIGH → DEF\nSET 저장" if kr else "POWER SET\nCH → LOW → HIGH → DEF\nSET Save",
        "rs485": "EXT:RS485\nBAUD → ADDR → TERM → PARITY\nSET 저장" if kr else "EXT:RS485\nBAUD → ADDR → TERM → PARITY\nSET Save",
        "mode_long": "MODE 길게" if kr else "MODE Long",
        "choose": "UP/DOWN 선택 후 SET" if kr else "Select by UP/DOWN, then SET",
        "common": "공통: MODE 짧게=필드 이동 / UP·DOWN=값 변경 / SET=저장 / SET 길게=선택 화면 복귀" if kr else "Common: MODE short=next field / UP-DOWN=change / SET=save / SET long=return",
        "screens_title": "LCD 화면 예시" if kr else "LCD Screen Examples",
        "stop": "기본 정지 화면" if kr else "Basic Stop Screen",
        "freq_ch": "FREQ 채널 선택" if kr else "FREQ Channel Select",
        "manual_freq": "FREQ 수동 설정" if kr else "Manual FREQ Setting",
        "power_set": "POWER 설정" if kr else "POWER Setting",
        "rs485_set": "RS485 설정" if kr else "RS485 Setting",
        "alarm": "알람 화면" if kr else "Alarm Screen",
        "remote_title": "REMOTE / EXT 운전 개념" if kr else "REMOTE / EXT Operation Concept",
        "remote_box": "REMOTE 운전\n외부 접점 ON/OFF\nREMOTE 입력 SHORT=RUN" if kr else "REMOTE Operation\nExternal Contact ON/OFF\nREMOTE SHORT=RUN",
        "ext_box": "EXT 운전\nRS-485 Modbus RTU\n마스터에서 읽기/쓰기" if kr else "EXT Operation\nRS-485 Modbus RTU\nRead/Write by Master",
        "remote_note": "REMOTE: 현장 접점 제어가 필요한 경우 사용" if kr else "REMOTE: Used for field contact control",
        "ext_note": "EXT: PLC/PC에서 RS-485 통신으로 상태 확인 및 설정 제어" if kr else "EXT: PLC/PC monitors and controls through RS-485",
        "warn_note": "주의: 통신 설정은 패널의 EXT:RS485 메뉴에서 저장" if kr else "Note: Save communication settings in EXT:RS485 menu",
        "wiring_title": "후면 연결 개념" if kr else "Rear Connection Concept",
        "power_in": "전원 입력\nAC 전원/접지" if kr else "Power Input\nAC Power / Earth",
        "output": "출력 커넥터\nNozzle / PZT" if kr else "Output Connector\nNozzle / PZT",
        "control": "Control I/O\nREMOTE/SENSOR/ALARM",
        "comm": "RS-485\nA/B + GND",
        "earth": "보호 접지\nPE 확인" if kr else "Protective Earth\nCheck PE",
        "wiring_note": "설치 전 출력선/센서선/접지/RS-485 A,B 극성을 반드시 확인한다." if kr else "Before installation, check output line, sensor line, earth and RS-485 A/B polarity.",
    }

    figures = {}

    img, d, path = canvas(f"{prefix}_panel_layout.png", text["panel_title"])
    d.rounded_rectangle((90, 115, 1110, 650), radius=30, fill=(245, 248, 252), outline=(80, 110, 150), width=4)
    d.rounded_rectangle((300, 170, 900, 330), radius=15, fill=(212, 238, 220), outline=(40, 100, 60), width=4)
    d.rectangle((335, 205, 865, 295), fill=(28, 55, 38))
    d.text((380, 218), "0.00W     0000kHz", font=FONT_LCD, fill=(180, 255, 180))
    d.text((380, 255), "NORMAL STOP PWR1", font=FONT_LCD, fill=(180, 255, 180))
    d.text((520, 135), text["lcd"], font=FONT_HEAD, fill=(25, 65, 110))
    for i, label in enumerate(["NORMAL", "H/L SET", "8 POWER", "REMOTE", "EXT", "TX/RX"]):
        x = 150 + i * 170
        d.ellipse((x, 390, x + 38, 428), fill=(80, 180, 95), outline=(30, 90, 40), width=2)
        d.text((x - 28, 440), label, font=FONT_SMALL, fill=(40, 40, 40))
    buttons = [("START/STOP", "RUN/STOP"), ("MODE", text["mode"]), ("UP", text["up"]), ("DOWN", text["down"]), ("SET", text["set"])]
    for i, (name, desc) in enumerate(buttons):
        x = 135 + i * 210
        d.rounded_rectangle((x, 535, x + 150, 610), radius=18, fill=(232, 242, 252), outline=(30, 80, 130), width=3)
        draw_center(d, (x, 535, x + 150, 585), name, FONT_SMALL, fill=(20, 55, 95))
        draw_center(d, (x, 580, x + 150, 610), desc, FONT_SMALL, fill=(80, 80, 80))
    img.save(path)
    figures["panel"] = path

    img, d, path = canvas(f"{prefix}_menu_flow.png", text["flow_title"], (1200, 820))
    boxes = [
        (70, 130, 360, 250, text["select"]),
        (455, 130, 745, 250, text["setting"]),
        (70, 390, 360, 560, text["freq"]),
        (455, 390, 745, 560, text["power"]),
        (840, 390, 1130, 560, text["rs485"]),
    ]
    for x1, y1, x2, y2, label in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=(240, 248, 255), outline=(35, 85, 135), width=3)
        draw_center(d, (x1, y1, x2, y2), label, FONT, fill=(30, 30, 30))
    arrow(d, (360, 190), (455, 190))
    d.text((375, 155), text["mode_long"], font=FONT_SMALL, fill=(35, 85, 135))
    arrow(d, (600, 250), (210, 390))
    arrow(d, (600, 250), (600, 390))
    arrow(d, (600, 250), (985, 390))
    d.text((370, 310), text["choose"], font=FONT, fill=(35, 85, 135))
    d.text((85, 640), text["common"], font=FONT, fill=(50, 50, 50))
    img.save(path)
    figures["flow"] = path

    img, d, path = canvas(f"{prefix}_screen_examples.png", text["screens_title"], (1200, 780))
    examples = [
        (text["stop"], "0.00W     0000kHz", "NORMAL STOP PWR1"),
        (text["freq_ch"], "01CH      0400kHz", "CH01 E0000 A0000"),
        (text["manual_freq"], "START(3SEC):AUTO", "FREQ:0406kHz"),
        (text["power_set"], "P-SET P1", "DEF:0.50W"),
        (text["rs485_set"], "EXT:RS485 SET", "BAUD:9600"),
        (text["alarm"], "   ALARM STOP", "ERR1 LOW ALARM"),
    ]
    for i, (title, line1, line2) in enumerate(examples):
        col = i % 2
        row = i // 2
        x = 90 + col * 560
        y = 120 + row * 195
        d.text((x, y), title, font=FONT_HEAD, fill=(25, 65, 110))
        d.rounded_rectangle((x, y + 38, x + 470, y + 155), radius=12, fill=(28, 55, 38), outline=(60, 130, 80), width=4)
        d.text((x + 35, y + 58), line1, font=FONT_LCD, fill=(180, 255, 180))
        d.text((x + 35, y + 103), line2, font=FONT_LCD, fill=(180, 255, 180))
    img.save(path)
    figures["screens"] = path

    img, d, path = canvas(f"{prefix}_remote_ext.png", text["remote_title"], (1200, 720))
    d.rounded_rectangle((90, 150, 395, 330), radius=18, fill=(255, 246, 222), outline=(150, 105, 30), width=3)
    draw_center(d, (90, 150, 395, 330), text["remote_box"], FONT)
    d.rounded_rectangle((805, 150, 1110, 330), radius=18, fill=(226, 239, 218), outline=(50, 110, 50), width=3)
    draw_center(d, (805, 150, 1110, 330), text["ext_box"], FONT)
    d.rounded_rectangle((465, 230, 735, 450), radius=18, fill=(232, 242, 252), outline=(35, 85, 135), width=3)
    draw_center(d, (465, 230, 735, 450), "CSF JET\nMulti Megasonic\nGenerator", FONT_HEAD, fill=(25, 65, 110))
    arrow(d, (395, 240), (465, 300))
    arrow(d, (805, 240), (735, 300))
    d.text((130, 390), text["remote_note"], font=FONT, fill=(60, 60, 60))
    d.text((130, 430), text["ext_note"], font=FONT, fill=(60, 60, 60))
    d.text((130, 500), text["warn_note"], font=FONT, fill=(180, 45, 45))
    img.save(path)
    figures["remote"] = path

    img, d, path = canvas(f"{prefix}_wiring.png", text["wiring_title"], (1200, 760))
    for x1, y1, x2, y2, label in [
        (80, 150, 330, 290, text["power_in"]),
        (475, 150, 725, 290, text["output"]),
        (870, 150, 1120, 290, text["control"]),
        (250, 440, 520, 590, text["comm"]),
        (680, 440, 950, 590, text["earth"]),
    ]:
        d.rounded_rectangle((x1, y1, x2, y2), radius=16, fill=(245, 248, 252), outline=(35, 85, 135), width=3)
        draw_center(d, (x1, y1, x2, y2), label, FONT)
    arrow(d, (330, 220), (475, 220))
    arrow(d, (725, 220), (870, 220))
    arrow(d, (520, 515), (680, 515))
    d.text((110, 660), text["wiring_note"], font=FONT, fill=(180, 45, 45))
    img.save(path)
    figures["wiring"] = path

    return figures


def md_table(headers, rows):
    out = "| " + " | ".join(headers) + " |\n"
    out += "| " + " | ".join(["---"] * len(headers)) + " |\n"
    for row in rows:
        out += "| " + " | ".join(str(v) for v in row) + " |\n"
    return out


def freq_rows(lang: str):
    label = "주파수 프로파일" if lang == "KR" else "Frequency profile"
    return [[f"CH{i:02d}", f"{400 + (i - 1) * 200}kHz", f"{label} {i}"] for i in range(1, 11)]


def content(lang: str, figures):
    kr = lang == "KR"
    if kr:
        return {
            "title": "LCD 패널 사용자 운전 매뉴얼",
            "version": "v3.2-KR",
            "purpose": "사용자가 본 문서만 보고 설치, 기본 운전, 메뉴 설정, REMOTE/EXT 운전, 알람 조치를 수행할 수 있도록 한다.",
            "fig": {
                "panel": "LCD 조작 패널 구성",
                "screens": "LCD 화면 예시",
                "wiring": "후면 연결 개념",
                "flow": "LCD 메뉴 흐름도",
                "remote": "REMOTE / EXT 운전 개념",
            },
            "summary": [["발진 시작/정지", "선택 화면에서 START/STOP"], ["선택 항목 이동", "MODE 짧게: MODE -> FREQ CH -> 8 POWER"], ["설정 메뉴 진입", "선택 화면에서 MODE 길게"], ["값 변경", "UP/DOWN, 길게 누르면 빠르게 연속 변경"], ["저장", "SET 짧게, 성공 시 비프 2회"], ["선택 화면 복귀", "설정 화면에서 SET 약 2초"], ["Auto Tuning", "FREQ SET에서 START/STOP 약 3초"], ["알람 해제", "원인 확인 후 알람 화면에서 SET"]],
            "summary_note": "비프음: 일반 버튼 1회, 길게 누름 인식 긴 1회, 저장 성공 2회, 저장 실패 3회, 알람 긴 1회.",
            "spec": [["모델명", "CSF JET Multi Megasonic Generator"], ["표시 장치", "LCD1602, 2행 x 16문자"], ["버튼", "START/STOP, MODE, UP, DOWN, SET"], ["상태 LED", "NORMAL, H/L SET, 8 POWER, REMOTE, EXT, TX/RX"], ["주파수", "10CH, 기본 400kHz~2200kHz"], ["주파수 설정", "수동 설정 및 Auto Tuning"], ["LC BIT", "16단계, 0000~1111"], ["POWER", "8개 프로파일, LOW/HIGH/DEF"], ["출력 설정", "0.20W~5.00W"], ["통신", "RS-485 Half-Duplex, Modbus RTU"], ["기본 통신", "ADDR 1, 9600bps, NONE-8-1"], ["저장", "내부 FLASH 저장, 전원 OFF 후 유지"]],
            "install": ["진동 영향을 받지 않는 안정된 위치에 설치한다.", "산성/알칼리성 가스, 전도성 먼지, 유증기, 고온 환경을 피한다.", "방열을 위해 전후좌우 통풍 공간을 확보한다.", "보호접지(PE)를 먼저 연결한 후 전원을 인가한다.", "지정된 노즐/트랜스듀서와 케이블을 사용한다.", "출력 케이블 종류나 길이를 임의로 변경하지 않는다.", "배선 작업 전 반드시 전원을 끈다."],
            "use": ["출력 커넥터가 빠진 상태에서 운전하지 않는다.", "알람 발생 후 원인 확인 없이 반복 재기동하지 않는다.", "SENSOR가 OPEN이면 발진이 금지되거나 정지된다.", "설치 또는 노즐 교체 후에는 낮은 출력부터 단계적으로 운전한다.", "통신 운전 시 마스터와 장치의 주소, 속도, 패리티를 반드시 일치시킨다."],
            "parts": [["LCD", "출력, 주파수, 운전 모드, 설정 메뉴, 알람 표시"], ["START/STOP", "발진 시작/정지, FREQ SET에서 길게 누르면 Auto Tuning"], ["MODE", "선택 항목/편집 필드 이동, 길게 누르면 설정 진입 또는 상위 복귀"], ["UP/DOWN", "값 증가/감소, 길게 누르면 빠른 반복"], ["SET", "저장/확정, 설정 화면에서 길게 누르면 선택 화면 복귀"], ["LED", "현재 모드와 통신 상태 표시"]],
            "rear": [["전원", "입력 전원 규격 및 보호접지"], ["출력", "노즐/트랜스듀서 전용 케이블 연결"], ["SENSOR", "정상 상태 SHORT, OPEN 시 Err7/Err8"], ["REMOTE", "외부 접점 운전 입력"], ["RS-485", "A/B 극성, GND, 주소/속도/패리티"], ["ALARM", "GO, LOW, HIGH, TRANSDUCER, OPERATE 접점"]],
            "connect": ["전원을 끈다.", "출력 케이블을 지정 노즐/트랜스듀서에 연결한다.", "SENSOR 회로가 정상(SHORT)인지 확인한다.", "REMOTE 또는 ALARM 접점을 사용할 경우 제어 커넥터에 배선한다.", "EXT 운전을 사용할 경우 RS-485 A/B/GND를 연결한다.", "보호접지를 확인하고 전원을 인가한다."],
            "ready": ["공정 조건에 따라 노즐에 DI water 또는 액체가 공급되는지 확인한다.", "출력 케이블, SENSOR, REMOTE, RS-485 배선을 확인한다.", "전원 ON 후 LCD 표시와 LED 상태를 확인한다.", "필요한 FREQ 채널과 POWER 채널을 선택한다.", "최초 운전은 낮은 출력으로 시작한다."],
            "normal": ["전원 ON 후 NORMAL LED를 확인한다.", "필요하면 FREQ CH와 POWER 채널을 선택한다.", "START/STOP을 눌러 발진을 시작한다.", "LCD의 RUN 표시와 출력값을 확인한다.", "정지하려면 START/STOP을 다시 누른다."],
            "freq_select": "SELECT 화면에서 MODE로 FREQ CH 항목을 선택하고 UP/DOWN으로 채널을 고른 뒤 SET을 누른다.",
            "freq_set": ["MODE 길게 -> SETTING MODE 진입.", "1.FREQ SET 선택 후 SET.", "MODE 짧게로 CH -> FREQ -> BIT 필드를 이동한다.", "깜빡이는 값은 UP/DOWN으로 변경한다.", "SET을 눌러 현재 CH의 주파수와 BIT를 저장한다.", "저장 성공 시 비프음이 2회 울린다."],
            "bit_note": "BIT는 LC 릴레이 조합이며 0000~1111의 16단계이다. E는 채널에 저장된 기대 BIT, A는 현재 실제 적용 BIT이다.",
            "autotune": ["FREQ SET에서 대상 CH를 선택한다.", "START/STOP을 약 3초 누른다.", "LCD에 튜닝 진행 상태가 표시된다.", "성공 시 AUTO TUNE DONE이 표시된다.", "필요한 경우 SET을 눌러 저장한다."],
            "autotune_note": "실제 노즐/트랜스듀서 연결 상태에서 낮은 출력으로 수행한다. NO DETECTED 또는 TUNE TIMEOUT이면 배선, 부하, 노즐 상태를 점검한다.",
            "power_fields": [["CH", "POWER 프로파일 번호", "PWR1~PWR8"], ["LOW", "저출력 알람 기준", "0.10W~0.95W"], ["HIGH", "고출력 알람 기준", "0.20W~5.00W"], ["DEF", "기본 출력", "0.20W~5.00W"]],
            "power_note": "MODE 길게 -> 2.POWER SET -> SET -> MODE로 CH/LOW/HIGH/DEF 이동 -> UP/DOWN 변경 -> SET 저장.",
            "remote": ["발진을 정지하고 전원을 끈다.", "REMOTE 접점 배선을 연결한다.", "전원을 켠다.", "SELECT 화면에서 MODE 항목을 선택하고 UP/DOWN으로 REMOTE를 선택한다.", "SET을 눌러 저장한다.", "REMOTE 접점 SHORT 시 RUN, OPEN 시 STOP으로 동작한다."],
            "remote_note": "여러 대를 동시에 제어할 경우 각 장비는 독립 접점을 사용하는 것을 권장한다.",
            "ext": ["SELECT 화면에서 MODE 항목을 선택한다.", "UP/DOWN으로 EXT를 선택하고 SET 저장한다.", "MODE 길게 -> SETTING MODE -> 3.EXT:RS485 진입.", "BAUD, ADDR, TERM, PARITY를 설정한다.", "SET으로 저장한다.", "마스터 장비에서 동일한 통신 조건으로 접속한다."],
            "rs485_fields": [["BAUD", "9600/19200/38400/115200", "통신 속도"], ["ADDR", "1~247", "슬레이브 주소"], ["TERM", "OFF/ON", "종단 설정 기록, 하드웨어 제어 예약"], ["PARITY", "NONE-8-1/EVEN-8-1", "패널 설정 패리티"]],
            "sensor": [["SHORT", "정상"], ["정지 중 OPEN", "Err7 SENSOR OFF"], ["운전 중 OPEN", "Err8 SENSOR RUN"]],
            "remote_table": [["OPEN", "STOP"], ["SHORT", "RUN"]],
            "alarm": [["GO", "출력이 정상 범위"], ["LOW ALARM", "Err1 발생"], ["HIGH ALARM", "Err2 발생"], ["TRANSDUCER ALARM", "Err5 발생"], ["OPERATE", "통합 운전/알람 접점"]],
            "errors": [["Err1", "LOW ALARM", "부하/LOW 설정 확인 후 SET"], ["Err2", "HIGH ALARM", "부하/HIGH 설정 확인 후 SET"], ["Err5", "TRANSDUCER ALARM", "노즐/케이블/부하 확인 후 SET 또는 재기동"], ["Err6", "SETTING ERROR", "허용 범위 값으로 재입력"], ["Err7", "SENSOR OFF", "전원 OFF -> SENSOR 정상화 -> 전원 ON"], ["Err8", "SENSOR RUN", "전원 OFF -> SENSOR 정상화 -> 전원 ON"]],
            "comm": [["인터페이스", "RS-485 Half-Duplex"], ["프로토콜", "Modbus RTU"], ["기본 설정", "ADDR 1, 9600bps, NONE-8-1"], ["지원 속도", "9600/19200/38400/115200"], ["설정 메뉴", "SETTING MODE -> 3.EXT:RS485"], ["기능 코드", "0x03 / 0x06 / 0x10"]],
            "regs": [["0x0000", "R/W", "RUN 제어", "0=STOP, 1=RUN"], ["0x0001", "R/W", "주파수", "x0.1kHz"], ["0x0002", "R/W", "Duty", "x0.1%"], ["0x0006", "R/W", "운전 모드", "0=NORMAL,1=REMOTE,2=EXT"], ["0x0010", "R/W", "슬레이브 주소", "1~247"], ["0x0011", "R/W", "Baud 인덱스", "0~3"], ["0x0012", "R/W", "Parity", "0=None,1=Even,2=Odd"]],
            "storage": "SET으로 저장한 값은 내부 FLASH에 저장되어 전원 OFF 후에도 유지된다. 저장 대상은 운전 모드, FREQ 채널/주파수/BIT, POWER 채널/LOW/HIGH/DEF, RS-485 설정이다.",
            "trouble": [["출력 없음", "STOP/알람/SENSOR OPEN", "LCD와 알람, SENSOR 확인"], ["출력 후 정지", "LOW/HIGH 알람", "POWER LOW/HIGH 설정 확인"], ["전원 ON 시 Err7", "SENSOR OPEN", "전원 OFF 후 SENSOR SHORT 확인"], ["주파수가 다름", "FREQ CH 또는 저장 누락", "채널 선택 후 SET 저장"], ["RS-485 응답 없음", "A/B, 주소, 속도, 패리티 불일치", "배선과 EXT:RS485 설정 확인"], ["저장값 미복원", "SET 미저장 또는 저장 실패", "SET 후 비프 2회 확인"]],
            "maintenance": ["일일: 케이블/커넥터, LCD/LED, 알람 이력 확인", "주간: 출력 재현성, RS-485 응답, REMOTE 접점 확인", "월간: 트랜스듀서/케이블 절연, 알람 기준, 중요 설정 백업"],
            "as": "서비스 요청 시 모델명, 시리얼 번호, 알람 코드, 발생 조건, 설치 사진, 재현 절차를 함께 제공한다. 임의 개조, 오배선, 자연재해, 소모품은 보증 범위에서 제외될 수 있다.",
            "rev": "한글 사용자 운전 매뉴얼 분리 생성, LCD 조작/메뉴/REMOTE/EXT/알람/통신/점검 절차 정리",
        }

    return {
        "title": "LCD Panel User Operation Manual",
        "version": "v3.2-EN",
        "purpose": "This manual enables the user to install the generator, operate it from the LCD panel, configure menus, use REMOTE/EXT operation, and respond to alarms.",
        "fig": {
            "panel": "LCD Operation Panel",
            "screens": "LCD Screen Examples",
            "wiring": "Rear Connection Concept",
            "flow": "LCD Menu Flow",
            "remote": "REMOTE / EXT Operation Concept",
        },
        "summary": [["Run/Stop", "Press START/STOP on the select screen"], ["Move selected item", "Short press MODE: MODE -> FREQ CH -> 8 POWER"], ["Enter setting menu", "Long press MODE on the select screen"], ["Change value", "Press UP/DOWN; long press repeats quickly"], ["Save", "Short press SET; two beeps indicate success"], ["Return to select screen", "Long press SET for about 2 seconds in setting screens"], ["Auto Tuning", "Long press START/STOP for about 3 seconds in FREQ SET"], ["Clear alarm", "Check the cause, then press SET on the alarm screen"]],
        "summary_note": "Buzzer guide: normal key 1 beep, long-press recognition 1 longer beep, save success 2 beeps, save failure 3 beeps, alarm 1 long beep.",
        "spec": [["Model name", "CSF JET Multi Megasonic Generator"], ["Display", "LCD1602, 2 lines x 16 characters"], ["Buttons", "START/STOP, MODE, UP, DOWN, SET"], ["Status LEDs", "NORMAL, H/L SET, 8 POWER, REMOTE, EXT, TX/RX"], ["Frequency", "10 channels, default 400kHz to 2200kHz"], ["Frequency setting", "Manual setting and Auto Tuning"], ["LC BIT", "16 steps, 0000 to 1111"], ["POWER", "8 profiles, LOW/HIGH/DEF"], ["Output setting", "0.20W to 5.00W"], ["Communication", "RS-485 Half-Duplex, Modbus RTU"], ["Default communication", "ADDR 1, 9600bps, NONE-8-1"], ["Storage", "Internal FLASH, retained after power-off"]],
        "install": ["Install the generator on a stable support where vibration does not affect the unit.", "Avoid acidic or alkaline gas, conductive dust, oil mist, and high-temperature environments.", "Secure ventilation space around the unit.", "Connect protective earth before applying power.", "Use the assigned nozzle/transducer and matched cable.", "Do not change the output cable type or length without confirmation.", "Turn off power before wiring work."],
        "use": ["Do not operate while the output connector is disconnected.", "Do not repeatedly restart after an alarm without checking the cause.", "If SENSOR is open, oscillation is inhibited or stopped.", "After installation or nozzle replacement, start at low output and increase gradually.", "For communication operation, match the address, baudrate, and parity with the master."],
        "parts": [["LCD", "Displays output, frequency, operating mode, setting menu, and alarms"], ["START/STOP", "Starts/stops oscillation; long press in FREQ SET starts Auto Tuning"], ["MODE", "Moves selected item/edit field; long press enters setting or returns to upper menu"], ["UP/DOWN", "Increases/decreases values; long press repeats quickly"], ["SET", "Saves/confirms values; long press returns to select screen from setting screens"], ["LED", "Indicates mode and communication status"]],
        "rear": [["Power", "Input power rating and protective earth"], ["Output", "Dedicated nozzle/transducer cable connection"], ["SENSOR", "Normal state is SHORT; OPEN causes Err7/Err8"], ["REMOTE", "External contact operation input"], ["RS-485", "A/B polarity, GND, address/baud/parity"], ["ALARM", "GO, LOW, HIGH, TRANSDUCER, OPERATE contacts"]],
        "connect": ["Turn off power.", "Connect the output cable to the assigned nozzle/transducer.", "Confirm that the SENSOR circuit is normal (SHORT).", "Wire the control connector when using REMOTE or ALARM contacts.", "Wire RS-485 A/B/GND when using EXT operation.", "Check protective earth, then apply power."],
        "ready": ["Confirm DI water or process liquid is supplied to the nozzle when required.", "Check output cable, SENSOR, REMOTE, and RS-485 wiring.", "Turn on power and check the LCD and LEDs.", "Select the required FREQ channel and POWER channel.", "Start the first operation at low output."],
        "normal": ["Turn on power and check the NORMAL LED.", "Select the FREQ CH and POWER channel if needed.", "Press START/STOP to start oscillation.", "Check RUN display and output value on the LCD.", "Press START/STOP again to stop."],
        "freq_select": "On the SELECT screen, use MODE to select FREQ CH, use UP/DOWN to choose a channel, then press SET.",
        "freq_set": ["Long press MODE to enter SETTING MODE.", "Select 1.FREQ SET and press SET.", "Short press MODE to move CH -> FREQ -> BIT fields.", "Change the blinking value with UP/DOWN.", "Press SET to save the frequency and BIT of the current channel.", "Two beeps indicate successful save."],
        "bit_note": "BIT is the LC relay combination and has 16 steps from 0000 to 1111. E means expected BIT saved in the channel. A means the currently applied actual BIT.",
        "autotune": ["Select the target channel in FREQ SET.", "Press START/STOP for about 3 seconds.", "The LCD shows tuning progress.", "AUTO TUNE DONE appears when tuning succeeds.", "Press SET to save if needed."],
        "autotune_note": "Run Auto Tuning at low output with the actual nozzle/transducer connected. If NO DETECTED or TUNE TIMEOUT appears, check wiring, load, and nozzle condition.",
        "power_fields": [["CH", "POWER profile number", "PWR1 to PWR8"], ["LOW", "Low output alarm threshold", "0.10W to 0.95W"], ["HIGH", "High output alarm threshold", "0.20W to 5.00W"], ["DEF", "Default output", "0.20W to 5.00W"]],
        "power_note": "Long press MODE -> 2.POWER SET -> SET -> use MODE to move CH/LOW/HIGH/DEF -> change with UP/DOWN -> save with SET.",
        "remote": ["Stop oscillation and turn off power.", "Connect the REMOTE contact wiring.", "Turn on power.", "On the SELECT screen, select MODE and choose REMOTE with UP/DOWN.", "Press SET to save.", "REMOTE contact SHORT means RUN; OPEN means STOP."],
        "remote_note": "When controlling multiple generators, use an independent contact for each unit.",
        "ext": ["Select MODE on the SELECT screen.", "Choose EXT with UP/DOWN and press SET to save.", "Long press MODE -> SETTING MODE -> 3.EXT:RS485.", "Set BAUD, ADDR, TERM, and PARITY.", "Press SET to save.", "Connect from the master using the same communication settings."],
        "rs485_fields": [["BAUD", "9600/19200/38400/115200", "Communication speed"], ["ADDR", "1 to 247", "Slave address"], ["TERM", "OFF/ON", "Termination setting record; hardware control reserved"], ["PARITY", "NONE-8-1/EVEN-8-1", "Panel parity setting"]],
        "sensor": [["SHORT", "Normal"], ["OPEN while stopped", "Err7 SENSOR OFF"], ["OPEN while running", "Err8 SENSOR RUN"]],
        "remote_table": [["OPEN", "STOP"], ["SHORT", "RUN"]],
        "alarm": [["GO", "Output is within normal range"], ["LOW ALARM", "Err1 occurred"], ["HIGH ALARM", "Err2 occurred"], ["TRANSDUCER ALARM", "Err5 occurred"], ["OPERATE", "Integrated operation/alarm contact"]],
        "errors": [["Err1", "LOW ALARM", "Check load/LOW setting, then press SET"], ["Err2", "HIGH ALARM", "Check load/HIGH setting, then press SET"], ["Err5", "TRANSDUCER ALARM", "Check nozzle/cable/load, then press SET or restart"], ["Err6", "SETTING ERROR", "Enter a value within the allowed range"], ["Err7", "SENSOR OFF", "Power OFF -> restore SENSOR -> power ON"], ["Err8", "SENSOR RUN", "Power OFF -> restore SENSOR -> power ON"]],
        "comm": [["Interface", "RS-485 Half-Duplex"], ["Protocol", "Modbus RTU"], ["Default setting", "ADDR 1, 9600bps, NONE-8-1"], ["Supported baudrates", "9600/19200/38400/115200"], ["Setting menu", "SETTING MODE -> 3.EXT:RS485"], ["Function codes", "0x03 / 0x06 / 0x10"]],
        "regs": [["0x0000", "R/W", "RUN control", "0=STOP, 1=RUN"], ["0x0001", "R/W", "Frequency", "x0.1kHz"], ["0x0002", "R/W", "Duty", "x0.1%"], ["0x0006", "R/W", "Operating mode", "0=NORMAL,1=REMOTE,2=EXT"], ["0x0010", "R/W", "Slave address", "1 to 247"], ["0x0011", "R/W", "Baud index", "0 to 3"], ["0x0012", "R/W", "Parity", "0=None,1=Even,2=Odd"]],
        "storage": "Values saved by SET are stored in internal FLASH and retained after power-off. Stored items include operating mode, FREQ channel/frequency/BIT, POWER channel/LOW/HIGH/DEF, and RS-485 settings.",
        "trouble": [["No output", "STOP/alarm/SENSOR open", "Check LCD, alarm, and SENSOR"], ["Stops after output starts", "LOW/HIGH alarm", "Check POWER LOW/HIGH setting"], ["Err7 at power ON", "SENSOR open", "Power OFF and check SENSOR short"], ["Unexpected frequency", "Wrong FREQ CH or not saved", "Select channel and save with SET"], ["No RS-485 response", "A/B, address, baudrate, or parity mismatch", "Check wiring and EXT:RS485 setting"], ["Saved value not restored", "SET not pressed or save failed", "Save again and confirm two beeps"]],
        "maintenance": ["Daily: check cables/connectors, LCD/LEDs, and alarm history", "Weekly: check output repeatability, RS-485 response, and REMOTE contact", "Monthly: check transducer/cable insulation, alarm thresholds, and setting backup"],
        "as": "When requesting service, provide model name, serial number, alarm code, operating condition, installation photos, and reproduction steps. Unauthorized modification, incorrect wiring, natural disasters, and consumables may be excluded from warranty.",
        "rev": "English user operation manual separated; LCD operation, menus, REMOTE/EXT operation, alarms, communication, and maintenance procedures organized",
    }


def build_markdown(lang: str, data, figures):
    kr = lang == "KR"
    h = {
        "summary": "< OPERATING SUMMARY >",
        "spec": "1. Specification",
        "caution": "2. Caution on Product Use",
        "install": "2.1 설치 주의사항" if kr else "2.1 Caution on Installation",
        "use": "2.2 사용 주의사항" if kr else "2.2 Caution on Use",
        "parts": "3. Name and Function of Each Part",
        "panel": "3.1 Operation Panel",
        "screens": "3.2 LCD 화면 예시" if kr else "3.2 LCD Screen Examples",
        "back": "3.3 Back of Generator",
        "connect": "4. Connecting Method",
        "ready": "5. Operation Ready",
        "operating": "6. Operating Method",
        "menu": "6.1 Menu Structure",
        "normal": "6.2 NORMAL 운전" if kr else "6.2 NORMAL Operation",
        "freq_sel": "6.3 FREQ 채널 선택" if kr else "6.3 FREQ Channel Selection",
        "freq_set": "6.4 FREQ SET: 주파수/BIT 수동 설정" if kr else "6.4 FREQ SET: Manual Frequency/BIT Setting",
        "auto": "6.5 Auto Tuning",
        "power": "6.6 POWER SET",
        "remote": "6.7 REMOTE 운전" if kr else "6.7 REMOTE Operation",
        "ext": "6.8 EXT 운전" if kr else "6.8 EXT Operation",
        "wiring": "7. Cable Wiring Diagram",
        "sensor": "7.1 SENSOR",
        "remote_io": "7.2 REMOTE",
        "alarm": "7.3 ALARM 출력" if kr else "7.3 ALARM Output",
        "error": "8. ERROR Display",
        "comm": "9. Communication Specification",
        "storage": "10. Setting Storage",
        "trouble": "11. Troubleshooting",
        "maint": "12. Maintenance",
        "as": "13. A/S",
        "rev": "14. Revision History",
    }
    out = f"# CSF JET Multi Megasonic Generator\n\n## {data['title']}\n\n"
    out += f"- {'문서 버전' if kr else 'Document Version'}: {data['version']}\n"
    out += f"- {'개정일' if kr else 'Revision Date'}: 2026-05-19\n"
    out += f"- {'적용 장비' if kr else 'Applicable Product'}: CSF JET Multi Megasonic Generator (LCD Panel Type)\n"
    out += f"- {'목적' if kr else 'Purpose'}: {data['purpose']}\n\n---\n\n"
    out += f"## {h['summary']}\n\n" + md_table(["항목" if kr else "Item", "조작" if kr else "Operation"], data["summary"]) + "\n" + data["summary_note"] + "\n\n---\n\n"
    out += f"## {h['spec']}\n\n" + md_table(["No.", "항목" if kr else "Item", "사양" if kr else "Specification"], [[i + 1, *row] for i, row in enumerate(data["spec"])]) + "\n"
    out += "구성품은 Generator 본체, 고주파 출력 케이블, 제어 커넥터/하네스, RS-485 통신선(옵션)으로 구성된다.\n\n" if kr else "The product consists of the generator, high-frequency output cable, control connector/harness, and RS-485 communication cable when selected.\n\n"
    out += f"## {h['caution']}\n\n### {h['install']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["install"])) + "\n\n"
    out += f"### {h['use']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["use"])) + "\n\n"
    out += f"## {h['parts']}\n\n### {h['panel']}\n\n![{data['fig']['panel']}](./{figures['panel'].name})\n\n" + md_table(["부품" if kr else "Part", "기능" if kr else "Function"], data["parts"]) + "\n"
    out += f"### {h['screens']}\n\n![{data['fig']['screens']}](./{figures['screens'].name})\n\n"
    out += ("기본 화면, FREQ 설정, POWER 설정, RS-485 설정, 알람 화면을 LCD에서 직접 확인할 수 있다.\n\n" if kr else "The LCD shows basic status, FREQ setting, POWER setting, RS-485 setting, and alarm screens.\n\n")
    out += f"### {h['back']}\n\n![{data['fig']['wiring']}](./{figures['wiring'].name})\n\n" + md_table(["연결부" if kr else "Connection", "확인 내용" if kr else "Check Point"], data["rear"]) + "\n"
    out += f"## {h['connect']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["connect"])) + "\n\n"
    out += f"## {h['ready']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["ready"])) + "\n\n"
    out += f"## {h['operating']}\n\n### {h['menu']}\n\n![{data['fig']['flow']}](./{figures['flow'].name})\n\n```text\nSELECT Screen\n  MODE short: MODE -> FREQ CH -> 8 POWER\n  MODE long : Enter SETTING MODE\n\nSETTING MODE\n  1.FREQ SET\n  2.POWER SET\n  3.EXT:RS485\n```\n\n"
    out += f"### {h['normal']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["normal"])) + "\n\n"
    out += f"### {h['freq_sel']}\n\n" + md_table(["채널" if kr else "Channel", "기본 주파수" if kr else "Default Frequency", "설명" if kr else "Description"], freq_rows(lang)) + "\n" + data["freq_select"] + "\n\n"
    out += f"### {h['freq_set']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["freq_set"])) + "\n\n" + data["bit_note"] + "\n\n"
    out += f"### {h['auto']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["autotune"])) + "\n\n" + data["autotune_note"] + "\n\n"
    out += f"### {h['power']}\n\n" + md_table(["필드" if kr else "Field", "의미" if kr else "Meaning", "범위" if kr else "Range"], data["power_fields"]) + "\n" + data["power_note"] + "\n\n"
    out += f"### {h['remote']}\n\n![{data['fig']['remote']}](./{figures['remote'].name})\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["remote"])) + "\n\n" + data["remote_note"] + "\n\n"
    out += f"### {h['ext']}\n\n" + "\n".join(f"{i + 1}. {v}" for i, v in enumerate(data["ext"])) + "\n\n" + md_table(["필드" if kr else "Field", "설정값" if kr else "Setting", "설명" if kr else "Description"], data["rs485_fields"]) + "\n"
    out += f"## {h['wiring']}\n\n### {h['sensor']}\n\n" + md_table(["상태" if kr else "State", "동작" if kr else "Operation"], data["sensor"]) + "\n"
    out += f"### {h['remote_io']}\n\n" + md_table(["REMOTE 입력" if kr else "REMOTE Input", "동작" if kr else "Operation"], data["remote_table"]) + "\n"
    out += f"### {h['alarm']}\n\n" + md_table(["출력" if kr else "Output", "의미" if kr else "Meaning"], data["alarm"]) + "\n"
    out += f"## {h['error']}\n\n```text\n   ALARM STOP\nERR1 LOW ALARM\n```\n\n" + md_table(["코드" if kr else "Code", "의미" if kr else "Meaning", "조치" if kr else "Action"], data["errors"]) + "\n"
    out += f"## {h['comm']}\n\n" + md_table(["항목" if kr else "Item", "사양" if kr else "Specification"], data["comm"]) + "\n" + md_table(["주소" if kr else "Address", "R/W", "명칭" if kr else "Name", "값" if kr else "Value"], data["regs"]) + "\n"
    out += f"## {h['storage']}\n\n{data['storage']}\n\n"
    out += f"## {h['trouble']}\n\n" + md_table(["증상" if kr else "Symptom", "원인 후보" if kr else "Possible Cause", "조치" if kr else "Action"], data["trouble"]) + "\n"
    out += f"## {h['maint']}\n\n" + "\n".join(f"- {v}" for v in data["maintenance"]) + "\n\n"
    out += f"## {h['as']}\n\n{data['as']}\n\n"
    out += f"## {h['rev']}\n\n" + md_table(["Version", "Date", "Description"], [[data["version"], "2026-05-19", data["rev"]]])
    return out


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.name = "맑은 고딕"
                run.font.size = Pt(9)
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(value)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "맑은 고딕"
                    run.font.size = Pt(9)
    doc.add_paragraph()


def add_image(doc: Document, path: Path, caption: str):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(6.3))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def add_numbered(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_docx(lang: str, data, figures, out_path: Path):
    kr = lang == "KR"
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "맑은 고딕"
    doc.styles["Normal"].font.size = Pt(10)
    for level in range(1, 4):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Instruction Manual")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("----- CSF JET Multi Megasonic Generator -----")
    run.bold = True
    run.font.size = Pt(16)
    info = doc.add_paragraph(f"{data['title']} / {data['version']} / 2026-05-19")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_image(doc, figures["panel"], data["fig"]["panel"])
    doc.add_page_break()

    doc.add_heading("< OPERATING SUMMARY >", level=1)
    add_table(doc, ["항목" if kr else "Item", "조작" if kr else "Operation"], data["summary"])
    doc.add_paragraph(data["summary_note"])

    doc.add_heading("1. Specification", level=1)
    add_table(doc, ["No.", "항목" if kr else "Item", "사양" if kr else "Specification"], [[i + 1, *row] for i, row in enumerate(data["spec"])])

    doc.add_heading("2. Caution on Product Use", level=1)
    doc.add_heading("2.1 설치 주의사항" if kr else "2.1 Caution on Installation", level=2)
    add_numbered(doc, data["install"])
    doc.add_heading("2.2 사용 주의사항" if kr else "2.2 Caution on Use", level=2)
    add_numbered(doc, data["use"])

    doc.add_heading("3. Name and Function of Each Part", level=1)
    add_image(doc, figures["screens"], data["fig"]["screens"])
    add_table(doc, ["부품" if kr else "Part", "기능" if kr else "Function"], data["parts"])

    doc.add_heading("4. Connecting Method", level=1)
    add_image(doc, figures["wiring"], data["fig"]["wiring"])
    add_table(doc, ["연결" if kr else "Connection", "확인" if kr else "Check Point"], data["rear"])
    add_numbered(doc, data["connect"])

    doc.add_heading("5. Operation Ready", level=1)
    add_numbered(doc, data["ready"])

    doc.add_heading("6. Operating Method", level=1)
    add_image(doc, figures["flow"], data["fig"]["flow"])
    doc.add_heading("NORMAL", level=2)
    add_numbered(doc, data["normal"])
    doc.add_heading("FREQ SET", level=2)
    add_numbered(doc, data["freq_set"])
    doc.add_paragraph(data["bit_note"])
    doc.add_heading("POWER SET", level=2)
    add_table(doc, ["필드" if kr else "Field", "의미" if kr else "Meaning", "범위" if kr else "Range"], data["power_fields"])
    doc.add_paragraph(data["power_note"])
    doc.add_heading("REMOTE / EXT", level=2)
    add_image(doc, figures["remote"], data["fig"]["remote"])
    add_numbered(doc, data["remote"])
    add_numbered(doc, data["ext"])
    add_table(doc, ["필드" if kr else "Field", "설정값" if kr else "Setting", "설명" if kr else "Description"], data["rs485_fields"])

    doc.add_heading("7. Cable Wiring Diagram", level=1)
    add_table(doc, ["SENSOR", "Operation"], data["sensor"])
    add_table(doc, ["REMOTE", "Operation"], data["remote_table"])
    add_table(doc, ["ALARM", "Meaning"], data["alarm"])

    doc.add_heading("8. ERROR Display", level=1)
    add_table(doc, ["코드" if kr else "Code", "의미" if kr else "Meaning", "조치" if kr else "Action"], data["errors"])

    doc.add_heading("9. Communication Specification", level=1)
    add_table(doc, ["항목" if kr else "Item", "사양" if kr else "Specification"], data["comm"])
    add_table(doc, ["주소" if kr else "Address", "R/W", "명칭" if kr else "Name", "값" if kr else "Value"], data["regs"])

    doc.add_heading("10. Setting Storage", level=1)
    doc.add_paragraph(data["storage"])
    doc.add_heading("11. Troubleshooting", level=1)
    add_table(doc, ["증상" if kr else "Symptom", "원인 후보" if kr else "Possible Cause", "조치" if kr else "Action"], data["trouble"])
    doc.add_heading("12. Maintenance", level=1)
    for item in data["maintenance"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("13. A/S", level=1)
    doc.add_paragraph(data["as"])
    doc.add_heading("14. Revision History", level=1)
    add_table(doc, ["Version", "Date", "Description"], [[data["version"], "2026-05-19", data["rev"]]])
    doc.save(out_path)


def write_outputs(lang: str, md_path: Path, docx_path: Path):
    figures = make_figures(lang)
    data = content(lang, figures)
    md_path.write_text(build_markdown(lang, data, figures), encoding="utf-8")
    build_docx(lang, data, figures, docx_path)


def main():
    write_outputs("KR", DOCS / f"{BASE}_KR.md", DOCS / f"{BASE}_KR.docx")
    write_outputs("EN", DOCS / f"{BASE}_EN.md", DOCS / f"{BASE}_EN.docx")
    # Keep the original filename as the Korean default for compatibility.
    write_outputs("KR", DOCS / f"{BASE}.md", DOCS / f"{BASE}.docx")
    print(DOCS / f"{BASE}_KR.md")
    print(DOCS / f"{BASE}_KR.docx")
    print(DOCS / f"{BASE}_EN.md")
    print(DOCS / f"{BASE}_EN.docx")
    print(DOCS / f"{BASE}.md")
    print(DOCS / f"{BASE}.docx")


if __name__ == "__main__":
    main()
