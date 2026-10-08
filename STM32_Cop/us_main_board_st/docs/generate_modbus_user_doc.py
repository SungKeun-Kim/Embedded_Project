"""PLC Modbus RTU Markdown 가이드를 배포용 DOCX로 변환한다."""

from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DOC_DIR = Path(__file__).resolve().parent
SOURCE = DOC_DIR / "PLC_Modbus_RTU_사용자_가이드.md"
OUTPUT = DOC_DIR / "STM32G474CBT6_Modbus_RTU_PLC_사용자_가이드_v1.3.docx"
CANONICAL_OUTPUT = DOC_DIR / "STM32G474CBT6_Modbus_RTU_PLC_사용자_가이드.docx"


def set_run_font(run, name="맑은 고딕", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def clean_inline(text):
    text = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def add_text(paragraph, text, *, code=False, bold=False):
    run = paragraph.add_run(clean_inline(text))
    if code:
        set_run_font(run, "Consolas", 9)
    else:
        set_run_font(run, size=10, bold=bold)
    return run


def add_table(document, lines):
    rows = [[clean_inline(x.strip()) for x in line.strip().strip("|").split("|")]
            for line in lines]
    headers = rows[0]
    data_rows = rows[2:]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, value, bold=True)

    for values in data_rows:
        row = table.add_row()
        prevent_row_split(row)
        for index in range(len(headers)):
            value = values[index] if index < len(values) else ""
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            add_text(p, value)

    document.add_paragraph()


def add_code_block(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.left_indent = Cm(0.4)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    add_text(paragraph, "\n".join(lines), code=True)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    heading_sizes = {1: 20, 2: 15, 3: 12}
    heading_colors = {1: (31, 78, 121), 2: (47, 84, 150), 3: (55, 55, 55)}
    for level, size in heading_sizes.items():
        style = styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*heading_colors[level])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level > 1 else 4)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("STM32G474CBT6 · Modbus RTU PLC 사용자 가이드"),
                 size=8, color=(100, 100, 100))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Firmware v1.2.0  |  문서 v1.3  |  "), size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    i = 0
    title_count = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if (line.strip().startswith("|") and i + 1 < len(lines)
                and re.match(r"^\s*\|?[\s:|-]+\|", lines[i + 1])):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(document, table_lines)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                                   if level == 1 and title_count < 2
                                   else WD_ALIGN_PARAGRAPH.LEFT)
            add_text(paragraph, heading.group(2), bold=True)
            if level == 1:
                title_count += 1
            i += 1
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if bullet:
            p = document.add_paragraph(style="List Bullet")
            add_text(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if numbered:
            # Word의 List Number style은 서로 떨어진 목록도 이전 번호에 이어 붙인다.
            # Markdown에 적힌 번호를 그대로 출력하여 절마다 1부터 다시 시작한다.
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            add_text(p, f"{numbered.group(1)}. {numbered.group(2)}")
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        add_text(paragraph, line)
        i += 1

    document.core_properties.title = "STM32G474CBT6 Modbus RTU PLC 사용자 가이드"
    document.core_properties.subject = "RS-485 Modbus RTU 통신 및 Holding Register 사양"
    document.core_properties.author = "STM32G474CBT6 초음파 제어 보드 개발팀"
    document.core_properties.keywords = "STM32G474CBT6, RS-485, Modbus RTU, PLC"
    document.save(OUTPUT)
    shutil.copyfile(OUTPUT, CANONICAL_OUTPUT)
    print(OUTPUT)
    print(CANONICAL_OUTPUT)


if __name__ == "__main__":
    build_document()
