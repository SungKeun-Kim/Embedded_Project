#!/usr/bin/env python3
"""STM32CubeProgrammer 최초/업데이트 GUI 사용순서 DOCX 생성."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SCREENSHOT = DOCS / "CubeProgramer.jpg"
OPTION_BYTES_SCREENSHOT = DOCS / "CubeProgramer-1.jpg"
OUTPUT = DOCS / "STM32CubeProgrammer_최초용_업데이트용_사용순서_v1.1.docx"

NAVY = "12385B"
BLUE = "1F6EAA"
LIGHT_BLUE = "DCEAF7"
GREEN = "DDEFE2"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False,
                  color: str | None = None, size: int = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, WHITE, 9)
        set_cell_shading(table.rows[0].cells[index], NAVY)
        if widths:
            table.rows[0].cells[index].width = Cm(widths[index])
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            set_cell_text(cells[col_index], value, False, None, 9)
            if widths:
                cells[col_index].width = Cm(widths[col_index])
            if row_index % 2 == 1:
                set_cell_shading(cells[col_index], "F7F9FB")
    return table


def add_banner(doc, title: str, body: str, fill: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(2)


def add_step(doc, number: int, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(1.1)
    table.columns[1].width = Cm(16.2)
    left, right = table.rows[0].cells
    set_cell_shading(left, BLUE)
    set_cell_text(left, str(number), True, WHITE, 12)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    right.text = ""
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    p2 = right.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("페이지 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 25, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11, NAVY),
    ):
        style = doc.styles[name]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_header_footer(section) -> None:
    header = section.header.paragraphs[0]
    header.text = "STM32G474CBT6 초음파 발진기 | STM32CubeProgrammer 사용순서"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.text = "Firmware v1.2.0 | 문서 v1.1 | 2026-10-08     "
    add_page_number(footer)


def build() -> Path:
    for image_path in (SCREENSHOT, OPTION_BYTES_SCREENSHOT):
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    configure_styles(doc)
    add_header_footer(section)

    # 표지
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("STM32CubeProgrammer")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("최초용 Factory / 업데이트용 HEX 사용순서")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run(
        "STM32G474CBT6 초음파 발진기 메인보드\nFirmware v1.2.0 / 문서 v1.1"
    ).font.size = Pt(13)
    doc.add_paragraph()
    add_table(
        doc,
        ["구분", "사용 파일", "설정값 처리", "사용 시점"],
        [
            ["최초용", "US_MAIN_BOARD_FACTORY_V1.2.0.hex",
             "출하 기본값 기록", "새 MCU 최초 Writing / 전체 초기화"],
            ["업데이트용", "US_MAIN_BOARD_UPDATE_V1.2.0.hex",
             "기존 설정 유지", "사용 중인 보드의 Firmware 업데이트"],
        ],
        [2.0, 6.2, 3.6, 5.2],
    )
    doc.add_paragraph()
    add_banner(doc, "가장 중요한 구분",
               "업데이트용에서는 Full chip erase를 실행하지 않는다. "
               "최초용은 Option Byte 설정과 출하 기본값 기록을 함께 수행한다.", YELLOW)
    p = doc.add_paragraph("작성 기준: 사용자가 제공한 현재 STM32CubeProgrammer GUI 화면")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    doc.add_page_break()

    # 화면 구성
    doc.add_heading("1. 현재 화면 구성", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREENSHOT), width=Cm(17.5))
    p = doc.add_paragraph("그림 1. 사용자가 제공한 STM32CubeProgrammer Erasing & Programming 화면")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    add_table(
        doc,
        ["화면 위치", "표시 이름", "용도"],
        [
            ["오른쪽 상단", "ST-LINK / Connect", "ST-LINK로 대상 MCU 연결"],
            ["오른쪽", "ST-LINK configuration", "SWD, 주파수, Under reset, Hardware reset 설정"],
            ["왼쪽", "File path / Browse", "Factory 또는 Update HEX 선택"],
            ["왼쪽", "Verify programming", "Writing 후 Flash 검증"],
            ["왼쪽", "Run after programming", "완료 후 MCU Reset 및 실행"],
            ["가운데", "Erase flash memory", "최초 전체 초기화에만 Full chip erase 사용"],
            ["아래", "Log", "성공, 오류, Verify 결과 확인"],
        ],
        [3.0, 5.0, 9.0],
    )
    doc.add_page_break()

    # 공통 준비
    doc.add_heading("2. 공통 연결 준비", level=1)
    add_banner(doc, "보드 출력 정지",
               "Writing 전에 초음파 출력을 정지하고 전력 출력부가 안전한 상태인지 확인한다.", YELLOW)
    add_step(doc, 1, "ST-LINK와 보드 연결",
             "SWDIO, SWCLK, NRST, GND를 연결하고 보드에 정상 전원을 공급한다. "
             "오른쪽 Target voltage가 약 3.1~3.3 V인지 확인한다.")
    add_step(doc, 2, "오른쪽 상단에서 ST-LINK 선택",
             "접속 방식은 ST-LINK로 선택한다. Serial number에는 연결한 ST-LINK가 표시되어야 한다.")
    add_step(doc, 3, "ST-LINK configuration 설정",
             "Port=SWD, Mode=Under reset, Reset mode=Hardware reset, Speed=Reliable, "
             "Shared=Disabled를 사용한다. Frequency는 먼저 4000 kHz로 시도하고 연결 오류가 나면 "
             "1000 kHz, 그래도 실패하면 100 kHz로 낮춘다.")
    add_step(doc, 4, "Connect 클릭",
             "오른쪽 위 Connect를 누른다. 연결되면 Not connected 표시가 Connected로 바뀌고 "
             "Target information에 Device ID, Flash size, CPU 정보가 표시된다.")
    add_banner(doc, "화면의 Unable to get core ID 오류",
               "Frequency를 1000 또는 100 kHz로 낮추고 Under reset / Hardware reset을 확인한 뒤 "
               "보드 전원과 NRST·SWDIO·SWCLK·GND 배선을 점검한다.", RED)
    doc.add_heading("공통 Programming 화면 설정", level=2)
    add_table(
        doc,
        ["항목", "설정", "설명"],
        [
            ["Start address", "빈 칸", "HEX 파일이 자체 주소를 포함하므로 직접 입력하지 않음"],
            ["Skip flash erase before programming", "체크 해제", "파일에 포함된 Application sector를 정상 Erase 후 Writing"],
            ["Verify programming", "체크", "Writing 데이터 검증"],
            ["Run after programming", "체크", "완료 후 애플리케이션 실행"],
            ["Full Flash memory checksum", "선택 안 함", "일반 Writing에는 필요 없음"],
        ],
        [5.8, 3.2, 8.0],
    )

    # 최초용
    doc.add_heading("3. 최초용 Factory HEX 사용순서", level=1)
    add_banner(doc, "사용 파일", "US_MAIN_BOARD_FACTORY_V1.2.0.hex", GREEN)
    doc.add_paragraph(
        "Factory HEX는 Application과 출하 기본 설정을 함께 기록한다. 새 MCU의 최초 Writing에 "
        "사용하며, 사용 중인 보드를 출하 상태로 완전히 초기화할 때에도 사용할 수 있다."
    )
    add_step(doc, 1, "Option Bytes 화면 열기",
             "연결된 상태에서 왼쪽 세로 메뉴의 OB 아이콘을 누르고 User Configuration 항목을 연다.")
    add_step(doc, 2, "BOOT Option Byte 설정",
             "nSWBOOT0=0, nBOOT0=1로 설정한다. 다른 Option Byte는 변경하지 않는다. "
             "Apply를 누른 뒤 연결이 끊기면 다시 ST-LINK / Under reset으로 Connect한다.")
    add_banner(doc, "BOOT 설정 이유",
               "PB8은 UP 버튼과 BOOT0 기능을 함께 사용한다. 위 Option Byte가 적용되어야 PB8의 "
               "외부 Pull-up 때문에 System Memory Bootloader로 진입하는 현상을 막을 수 있다.", LIGHT_BLUE)
    doc.add_page_break()
    doc.add_heading("3.1 Option bytes 화면에서 설정하는 방법", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OPTION_BYTES_SCREENSHOT), width=Cm(17.5))
    p = doc.add_paragraph("그림 2. 현재 STM32CubeProgrammer Option bytes 화면")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    add_table(
        doc,
        ["화면 조작", "설정 방법"],
        [
            ["왼쪽 OB 아이콘", "Option bytes 화면으로 이동"],
            ["Detailed view", "세부 Option Byte 목록 표시"],
            ["nSWBOOT0", "목록을 스크롤하여 찾고 체크 해제: 값 0"],
            ["nBOOT0", "목록을 스크롤하여 찾고 체크: 값 1"],
            ["Apply", "두 값을 확인한 다음 적용"],
        ],
        [5.0, 12.0],
    )
    add_banner(doc, "체크박스 값 읽는 방법",
               "체크됨은 1, 체크 해제는 0을 의미한다. 따라서 nSWBOOT0는 체크 해제, "
               "nBOOT0는 체크해야 한다. Apply 후 연결이 끊기면 Under reset으로 다시 Connect한다.", YELLOW)
    doc.add_page_break()
    add_step(doc, 3, "Full chip erase 실행",
             "왼쪽 Erasing & Programming 화면으로 돌아간다. 가운데 Erase flash memory 탭에서 "
             "Full chip erase를 실행한다. 새 MCU에서는 생략해도 되지만 동일한 표준 절차를 권장한다. "
             "기존 보드를 출하 초기화할 때에는 반드시 실행한다.")
    add_step(doc, 4, "Factory HEX 선택",
             "File path 옆 Browse를 누르고 production 폴더의 "
             "US_MAIN_BOARD_FACTORY_V1.2.0.hex를 선택한다. Start address는 빈 칸으로 둔다.")
    add_step(doc, 5, "검증 옵션 확인",
             "Skip flash erase before programming=해제, Verify programming=체크, "
             "Run after programming=체크로 설정한다.")
    add_step(doc, 6, "Start Programming 실행",
             "Start Programming을 누르고 하단 Log에서 File download complete, "
             "Download verified successfully 또는 이에 해당하는 성공 문구를 확인한다.")
    add_step(doc, 7, "최초 부팅 확인",
             "LCD가 정상 부팅되는지 확인한다. 필요하면 전원을 껐다 켠 뒤 Supervisor 메뉴에서 "
             "제품별 통신 설정과 PL 범위를 다시 설정한다.")
    doc.add_heading("Factory 기본 설정", level=2)
    add_table(
        doc,
        ["항목", "기본값"],
        [
            ["운전시간", "10 MIN"],
            ["주파수 Band", "28.0 / 40.0 / 68.0 / 80.0 kHz"],
            ["Sweep 폭 / 속도", "500 Hz / 100 Hz"],
            ["Modbus", "Slave ID 1 / 9600 bps / 8-E-1"],
            ["RS485 종단저항", "UNMOUNTED"],
            ["PL 표시", "PERCENT"],
        ],
        [7.0, 10.0],
    )
    doc.add_page_break()

    # 업데이트용
    doc.add_heading("4. 업데이트용 HEX 사용순서", level=1)
    add_banner(doc, "사용 파일", "US_MAIN_BOARD_UPDATE_V1.2.0.hex", GREEN)
    add_banner(doc, "주의: Full chip erase 금지",
               "업데이트용은 기존 설정을 유지하기 위한 파일이다. 가운데 Full chip erase 버튼이나 "
               "Automatic Mode의 Full chip erase를 실행하면 저장 설정이 삭제된다.", RED)
    doc.add_paragraph(
        "업데이트 HEX에는 Application 영역만 들어 있으며 마지막 2 KB 설정 Page는 포함하지 않는다. "
        "따라서 정상적인 sector erase 방식으로 Writing하면 현재 설정을 유지한다."
    )
    add_step(doc, 1, "ST-LINK 연결",
             "공통 연결 준비와 같은 방법으로 ST-LINK / SWD / Under reset / Hardware reset으로 Connect한다.")
    add_step(doc, 2, "업데이트 HEX 선택",
             "Browse를 누르고 production 폴더의 US_MAIN_BOARD_UPDATE_V1.2.0.hex를 선택한다. "
             "Start address는 빈 칸으로 둔다.")
    add_step(doc, 3, "Erase 관련 항목 확인",
             "Full chip erase를 실행하지 않는다. Skip flash erase before programming은 체크하지 않는다. "
             "CubeProgrammer가 HEX에 포함된 Application sector만 자동 Erase하도록 둔다.")
    add_step(doc, 4, "검증과 실행 옵션 확인",
             "Verify programming과 Run after programming을 체크한다.")
    add_step(doc, 5, "Start Programming 실행",
             "Start Programming을 누르고 하단 Log에서 Download와 Verify 성공을 확인한다.")
    add_step(doc, 6, "설정 유지 확인",
             "Reset 후 LCD와 Supervisor 메뉴를 확인한다. 기존 주파수, Sweep 폭·속도, 운전시간, "
             "Slave ID, Baud rate, PL 범위와 종단저항 설정이 유지되어야 한다.")
    add_table(
        doc,
        ["업데이트 후 확인 항목", "정상 기준"],
        [
            ["LCD 부팅", "애플리케이션 정상 시작"],
            ["주파수 / Sweep", "업데이트 전 사용자 설정 유지"],
            ["통신", "기존 Slave ID와 Baud rate 유지"],
            ["운전 상태", "안전을 위해 출력 OFF에서 시작"],
            ["PLC Lock / RUN", "전원 재인가 후 해제 및 OFF, PLC가 다시 명령"],
        ],
        [7.0, 10.0],
    )
    doc.add_page_break()

    # Automatic Mode 반복 라이팅
    doc.add_heading("5. Automatic Mode로 보드 교체 반복 라이팅", level=1)
    add_banner(doc, "양산 시 권장 방법",
               "첫 번째 보드에서 Connect와 Start automatic mode를 한 번만 실행한다. 그다음부터는 "
               "완료된 보드를 분리하고 새 보드를 연결하면 자동으로 Writing을 반복한다.", GREEN)
    doc.add_paragraph(
        "Automatic Mode는 STM32CubeProgrammer가 보드 분리와 다음 보드 연결을 감지하여 "
        "Erase, Download, Verify, Reset/Run을 순서대로 반복하는 기능이다. PC에는 ST-LINK를 "
        "한 개만 연결하여 사용한다."
    )
    add_step(doc, 1, "첫 번째 보드 연결",
             "보드 전원과 ST-LINK를 연결하고 오른쪽 위 Connect를 한 번 누른다. "
             "Connected와 Target information을 확인한다.")
    add_step(doc, 2, "HEX와 공통 옵션 설정",
             "File path에서 목적에 맞는 Factory 또는 Update HEX를 선택한다. "
             "Skip flash erase before programming=체크 해제, Verify programming=체크, "
             "Run after programming=체크로 설정한다.")
    doc.add_heading("Automatic Mode 체크 항목", level=2)
    add_table(
        doc,
        ["항목", "최초용 Factory", "업데이트용"],
        [
            ["선택 HEX", "...FACTORY_V1.2.0.hex", "...UPDATE_V1.2.0.hex"],
            ["Full chip erase", "체크", "체크 해제"],
            ["Download file", "체크", "체크"],
            ["Option bytes commands", "체크", "체크 해제"],
            ["Option Byte 명령", "-ob nSWBOOT0=0 nBOOT0=1", "입력하지 않음"],
        ],
        [5.1, 6.0, 5.9],
    )
    add_banner(doc, "업데이트용 주의",
               "업데이트용 Automatic Mode에서는 Full chip erase를 체크하지 않는다. "
               "체크하면 주파수, Sweep, 운전시간과 통신 설정이 삭제된다.", RED)
    doc.add_page_break()

    doc.add_heading("5.1 Automatic Mode 실행과 보드 교체", level=2)
    add_step(doc, 3, "Start automatic mode 실행",
             "Automatic Mode 영역이 보이도록 화면을 아래로 스크롤한다. 설정을 다시 확인한 뒤 "
             "Start automatic mode를 누른다. 첫 번째 보드의 작업이 즉시 시작된다.")
    add_step(doc, 4, "성공 Log 확인",
             "Programming, Verify와 Reset/Run 성공을 확인한다. 이어서 "
             "Please disconnect device and connect the next 또는 Waiting for device가 표시된다.")
    add_step(doc, 5, "완료된 보드 분리",
             "성공 Log가 나온 뒤 보드 전원을 끄고 SWD 커넥터를 분리한다. STM32CubeProgrammer가 "
             "분리를 인식할 때까지 기다린다.")
    add_step(doc, 6, "다음 보드 연결",
             "다음 보드에 SWD 커넥터를 연결하고 전원을 켠다. 자동 감지 후 같은 설정으로 "
             "Writing이 시작되므로 Connect나 Start Programming을 다시 누르지 않는다.")
    add_step(doc, 7, "반복 작업 종료",
             "마지막 보드의 성공 Log를 확인한 뒤 Stop automatic mode 또는 Cancel을 눌러 종료한다.")
    add_banner(doc, "반복 순서",
               "첫 보드 Connect → Start automatic mode → 성공 확인 → 보드 전원 OFF·분리 → "
               "다음 보드 연결·전원 ON → 자동 Writing", LIGHT_BLUE)
    doc.add_heading("Automatic Mode 작업 전 체크리스트", level=2)
    for text in (
        "PC에 ST-LINK가 한 개만 연결되어 있다.",
        "Factory와 Update HEX 중 올바른 파일을 선택했다.",
        "Update 작업에서는 Full chip erase가 체크 해제되어 있다.",
        "Verify programming과 Run after programming이 체크되어 있다.",
        "성공 Log가 표시된 뒤에만 보드를 교체한다.",
    ):
        p = doc.add_paragraph(style=None)
        p.add_run("□ ").bold = True
        p.add_run(text)
    p = doc.add_paragraph(
        "참고: STMicroelectronics UM2237, STM32CubeProgrammer Automatic mode"
    )
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    # 비교/금지 사항
    doc.add_heading("6. 최초용과 업데이트용 비교", level=1)
    add_table(
        doc,
        ["항목", "최초용 Factory", "업데이트용"],
        [
            ["HEX 파일", "...FACTORY_V1.2.0.hex", "...UPDATE_V1.2.0.hex"],
            ["Application", "기록", "기록"],
            ["설정 Page", "출하 기본값 기록", "파일에 포함하지 않음"],
            ["Full chip erase", "권장 / 초기화 시 필수", "금지"],
            ["Option Byte", "nSWBOOT0=0, nBOOT0=1 확인", "이미 설정된 보드는 변경 불필요"],
            ["기존 설정", "초기화", "유지"],
            ["사용 대상", "새 MCU / 출하 초기화", "현장 Firmware 업데이트"],
        ],
        [4.4, 6.3, 6.3],
    )
    doc.add_heading("누르면 안 되는 항목", level=2)
    add_banner(doc, "업데이트 시 금지",
               "Erase flash memory의 Full chip erase, Automatic Mode의 Full chip erase, "
               "잘못된 Factory HEX 선택", RED)
    doc.add_heading("파일 선택 전 최종 확인", level=2)
    add_table(
        doc,
        ["질문", "예", "선택 파일"],
        [
            ["새 MCU에 처음 Writing하는가?", "예", "FACTORY"],
            ["모든 설정을 출하값으로 초기화하는가?", "예", "FACTORY"],
            ["기존 설정을 유지하며 Firmware만 바꾸는가?", "예", "UPDATE"],
        ],
        [9.0, 2.0, 6.0],
    )
    doc.add_page_break()

    # 문제 해결
    doc.add_heading("7. 오류 해결", level=1)
    add_table(
        doc,
        ["증상 / Log", "확인 순서"],
        [
            ["Unable to get core ID",
             "Frequency 1000→100 kHz로 낮춤 → Under reset 확인 → Hardware reset 확인 → NRST/SWD 배선 점검"],
            ["Not connected",
             "ST-LINK 선택 → Serial number 확인 → 보드 전원 확인 → Connect 클릭"],
            ["Target voltage가 0 V",
             "보드 3.3 V 전원과 ST-LINK GND 공통 연결 확인"],
            ["Download 실패",
             "다른 Debug 프로그램 종료 → 재Connect → 낮은 SWD Frequency로 재시도"],
            ["Verify 실패",
             "전원 안정성 확인 → 100 kHz로 낮춤 → 케이블 단축 → 다시 Writing"],
            ["Writing 후 LCD가 실행되지 않음",
             "Run after programming 확인 → Reset → Factory 최초 작업이면 BOOT Option Byte 확인"],
            ["업데이트 후 설정이 초기화됨",
             "Full chip erase 또는 Factory HEX 사용 여부 확인. 사용자 설정을 다시 입력"],
            ["Automatic Mode가 시작되지 않음",
             "첫 번째 보드 Connect 확인 → Download file 체크 확인 → HEX 경로 확인"],
            ["다음 보드를 감지하지 못함",
             "Waiting for device 확인 → 이전 보드 분리 인식 대기 → 새 보드 전원과 SWD 배선 확인"],
            ["More than one ST-LINK probe detected",
             "PC에서 작업에 사용하지 않는 ST-LINK를 분리하고 한 개만 연결"],
        ],
        [6.2, 10.8],
    )
    doc.add_heading("성공 Log 확인", level=2)
    add_banner(doc, "정상 완료 기준",
               "Programming 완료, Verify 성공, MCU Reset/Run 성공이 표시되고 오류 문구가 없어야 한다.", GREEN)
    doc.add_heading("최종 체크리스트", level=2)
    for text in (
        "파일명이 FACTORY인지 UPDATE인지 확인했다.",
        "Verify programming을 체크했다.",
        "업데이트 작업에서 Full chip erase를 사용하지 않았다.",
        "하단 Log에서 Verify 성공을 확인했다.",
        "Reset 후 LCD와 저장 설정을 확인했다.",
    ):
        p = doc.add_paragraph(style=None)
        p.add_run("□ ").bold = True
        p.add_run(text)

    doc.add_paragraph()
    add_banner(doc, "배포 파일 위치",
               "프로젝트 production 폴더에 Factory HEX, Update HEX, 본 안내서를 함께 보관한다.", LIGHT_BLUE)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
